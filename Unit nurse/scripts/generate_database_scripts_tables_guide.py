from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DOCX = DOCS / "Plasmit-HMS-Database-Scripts-and-Tables-High-Level-Guide.docx"
MD = DOCS / "plasmit-hms-database-scripts-and-tables-high-level-guide.md"


def rgb(hex_color: str) -> tuple[int, int, int]:
    hex_color = hex_color.replace("#", "")
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border(cell, color: str = "CBD5E1") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        item = borders.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), color)


def heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Aptos Display"
        run.font.bold = True
        run.font.color.rgb = RGBColor(*rgb("14325C" if level <= 2 else "2563EB"))


def para(doc: Document, text: str, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10.2)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*rgb("334155"))


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(item)
        run.font.name = "Aptos"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(*rgb("334155"))


def code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, "0F172A")
    border(cell, "1E293B")
    paragraph = cell.paragraphs[0]
    for line in text.strip().splitlines():
        run = paragraph.add_run(line + "\n")
        run.font.name = "Consolas"
        run.font.size = Pt(8.2)
        run.font.color.rgb = RGBColor(*rgb("E2E8F0"))
    doc.add_paragraph()


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = tbl.cell(0, index)
        shade(cell, "14325C")
        border(cell)
        run = cell.paragraphs[0].add_run(header)
        run.font.name = "Aptos"
        run.font.size = Pt(8.2)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            border(cell)
            if row_index % 2 == 0:
                shade(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Aptos"
            run.font.size = Pt(8.0)
            run.font.color.rgb = RGBColor(*rgb("334155"))
    doc.add_paragraph()


def callout(doc: Document, title: str, items: list[str]) -> None:
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    shade(cell, "EAF4FF")
    border(cell, "BBD4F8")
    run = cell.paragraphs[0].add_run(title)
    run.font.name = "Aptos Display"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*rgb("14325C"))
    for item in items:
        paragraph = cell.add_paragraph(style="List Bullet")
        bullet_run = paragraph.add_run(item)
        bullet_run.font.name = "Aptos"
        bullet_run.font.size = Pt(9.4)
        bullet_run.font.color.rgb = RGBColor(*rgb("334155"))
    doc.add_paragraph()


SCRIPT_ORDER = [
    ["01", "create_database", "Creates database/schema", "plasmit_hms", "Run first in MySQL Workbench or CLI."],
    ["02", "create_foundation_tables", "Creates SaaS root tables", "tenants, hospitals, branches", "Must run before every other module."],
    ["03", "create_user_role_permission_tables", "Creates access control", "users, roles, permissions", "Required before API security and audit."],
    ["04", "create_patient_tables", "Creates patient/MPI tables", "patients, identifiers, contacts", "Patient is hospital-level."],
    ["05", "create_appointment_tables", "Creates scheduling tables", "appointments, slots, queues", "Appointment is branch-level."],
    ["06", "create_encounter_clinical_tables", "Creates clinical care tables", "encounters, vitals, diagnoses", "Encounter is branch-level clinical anchor."],
    ["07", "create_lab_radiology_tables", "Creates diagnostics tables", "lab_orders, results, radiology_reports", "Connects through encounter."],
    ["08", "create_pharmacy_tables", "Creates medicine and stock tables", "medicine_master, batches, stock, sales", "Stock is branch-level."],
    ["09", "create_billing_tables", "Creates revenue cycle tables", "invoices, invoice_items, payments", "Billing is branch-level."],
    ["10", "create_fhir_hl7_tables", "Creates integration tables", "fhir_resource_mapping, hl7_message_log", "Run after core business tables."],
    ["11", "create_reporting_summary_tables", "Creates dashboard summary tables", "daily summaries, snapshots", "Generated from transaction tables."],
    ["12", "seed_master_data", "Loads default master data", "statuses, roles, code masters", "Use repeatable migration where possible."],
]


FIRST_TABLES = [
    ["1", "tenants", "Root SaaS customer / hospital group", "None", "First table in the system."],
    ["2", "hospitals", "Hospital under tenant", "tenants.id", "One tenant can own many hospitals."],
    ["3", "branches", "Branch/facility/location under hospital", "hospitals.id", "One hospital can have many branches."],
    ["4", "departments", "Hospital/branch departments", "hospitals.id, branches.id optional", "OPD, IPD, ER, Lab, Pharmacy, Billing."],
    ["5", "healthcare_services", "Services provided by departments", "departments.id", "Maps to FHIR HealthcareService."],
    ["6", "users", "Doctors/staff/admin users", "tenants.id, hospitals.id", "Users get access using branch permissions."],
    ["7", "roles", "User role master", "tenant_id/hospital_id if custom", "Admin, doctor, nurse, lab, pharmacy, billing."],
    ["8", "permissions", "System permission master", "None or tenant_id", "Controls actions."],
    ["9", "role_permissions", "Role-permission mapping", "roles.id, permissions.id", "RBAC mapping."],
    ["10", "user_branch_access", "User branch access mapping", "users.id, branches.id", "Branch isolation starts here."],
]


MODULE_ORDER = [
    ["Foundation", "tenants -> hospitals -> branches -> departments -> services", "Root hierarchy and setup"],
    ["Access Control", "users -> roles -> permissions -> user_branch_access", "Login, RBAC and branch access"],
    ["Patient", "patients -> identifiers -> addresses -> contacts", "Hospital-level patient identity"],
    ["Appointment", "slots -> appointments -> queue/status", "Branch-level scheduling"],
    ["Encounter/Clinical", "encounters -> vitals -> diagnoses -> notes -> prescriptions", "Clinical source of truth"],
    ["Diagnostics", "lab/radiology orders -> samples/results/reports", "Order-to-result workflow"],
    ["Pharmacy", "medicine -> batch -> stock -> dispensing/sales", "Stock and medicine workflow"],
    ["Billing", "invoice -> invoice_items -> payments/refunds/claims", "Revenue cycle"],
    ["Integration", "FHIR mapping -> HL7 log -> external identifiers", "FHIR/HL7 readiness"],
    ["Reporting", "summary tables and snapshots", "Dashboard performance"],
]


def build_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    cover = doc.add_table(rows=1, cols=1)
    cell = cover.cell(0, 0)
    shade(cell, "14325C")
    border(cell, "14325C")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Plasmit HMS Database Scripts and Tables")
    run.font.name = "Aptos Display"
    run.font.size = Pt(21)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("High-Level Creation Order Guide")
    run2.font.name = "Aptos Display"
    run2.font.size = Pt(15)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(*rgb("D8E8FF"))
    doc.add_paragraph()

    table(doc, ["Project", "Database", "Technology", "Purpose"], [[
        "Plasmit Global Hospital Management System",
        "plasmit_hms",
        "MySQL + Spring Boot + JDBC + Flyway",
        "Explain which scripts and tables should be created first and why.",
    ]])

    heading(doc, "1. Executive Summary", 1)
    para(doc, "This document explains the high-level database script and table creation order for Plasmit HMS. The goal is to keep database creation clean, repeatable and safe for multi-developer work.")
    callout(doc, "Simple meaning", [
        "Database script means a .sql file that creates or changes database objects.",
        "Table means the real structure inside MySQL where data is stored.",
        "Migration means a versioned database script managed by Flyway.",
        "The first database should be plasmit_hms.",
        "The first table should be tenants.",
    ])

    heading(doc, "2. Database Name", 1)
    code(doc, """
CREATE DATABASE IF NOT EXISTS plasmit_hms
  CHARACTER SET utf8mb4
  COLLATE utf8mb4_0900_ai_ci;

USE plasmit_hms;
""")

    heading(doc, "3. What Is a Database Script?", 1)
    bullets(doc, [
        "A database script is a SQL file.",
        "It can create a database, create tables, add columns, create indexes, create seed data, create triggers or stored procedures.",
        "In a team project, developers should not manually change database directly.",
        "Every change should be written inside a script and committed to Git.",
        "For Spring Boot production projects, Flyway migration scripts are recommended.",
    ])

    heading(doc, "4. What Is a Table?", 1)
    bullets(doc, [
        "A table stores one business entity or transaction.",
        "Example: tenants stores SaaS customers or hospital groups.",
        "Example: hospitals stores hospitals under tenants.",
        "Example: patients stores patient identity data.",
        "Tables connect using primary keys and foreign keys.",
    ])

    heading(doc, "5. High-Level Script Order", 1)
    table(doc, ["Order", "Script Name", "Purpose", "Main Objects", "Note"], SCRIPT_ORDER)

    heading(doc, "6. First Tables to Create", 1)
    table(doc, ["Order", "Table", "Purpose", "Depends On", "Why Important"], FIRST_TABLES)

    heading(doc, "7. Flyway Naming Convention", 1)
    para(doc, "Use timestamp-based Flyway migration names so multiple developers do not create the same migration version.")
    code(doc, """
V202606030001__create_foundation_tables.sql
V202606030002__create_user_role_permission_tables.sql
V202606030003__create_patient_tables.sql
V202606030004__create_appointment_tables.sql
V202606030005__create_encounter_clinical_tables.sql
V202606030006__create_lab_radiology_tables.sql
V202606030007__create_pharmacy_tables.sql
V202606030008__create_billing_tables.sql
V202606030009__create_fhir_hl7_tables.sql
R__seed_default_master_data.sql
""")

    heading(doc, "8. Module-Wise Creation Flow", 1)
    table(doc, ["Module", "Table Flow", "Business Meaning"], MODULE_ORDER)

    heading(doc, "9. Minimum Tenant Isolation Columns", 1)
    table(doc, ["Table Type", "Required Columns", "Example"], [
        ["Global master", "id, code, name", "country_master, gender_master"],
        ["Tenant-level", "id, tenant_id", "tenant_settings"],
        ["Hospital-level", "id, tenant_id, hospital_id", "patients, departments"],
        ["Branch-level", "id, tenant_id, hospital_id, branch_id", "appointments, encounters, lab_orders, billing_invoices, pharmacy_stock"],
        ["Encounter-level", "id, tenant_id, hospital_id, branch_id, patient_id, encounter_id", "vitals, diagnoses, prescriptions, lab_results"],
    ])

    heading(doc, "10. Standard Query Pattern", 1)
    code(doc, """
SELECT *
FROM patients
WHERE tenant_id = :tenantId
  AND hospital_id = :hospitalId
  AND is_deleted = 0;

SELECT *
FROM appointments
WHERE tenant_id = :tenantId
  AND hospital_id = :hospitalId
  AND branch_id IN (:allowedBranchIds)
  AND is_deleted = 0;
""")

    heading(doc, "11. Final Recommendation", 1)
    callout(doc, "Recommended starting point", [
        "Create database: plasmit_hms.",
        "Run foundation script first.",
        "Create tenants table first.",
        "After tenants, create hospitals, branches, departments and users.",
        "Only after foundation tables, create patient, appointment, encounter, lab, pharmacy and billing tables.",
        "Use Flyway migrations for all future table changes.",
    ])

    return doc


def build_markdown() -> str:
    def md_table(headers: list[str], rows: list[list[str]]) -> list[str]:
        output = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join(["---"] * len(headers)) + " |",
        ]
        for row in rows:
            output.append("| " + " | ".join(cell.replace("|", "/") for cell in row) + " |")
        return output

    lines = [
        "# Plasmit HMS Database Scripts and Tables - High-Level Creation Order Guide",
        "",
        "## Executive Summary",
        "",
        "This document explains what database scripts and tables mean, which script should run first, and which tables should be created first for Plasmit HMS.",
        "",
        "- Database name: `plasmit_hms`",
        "- First table: `tenants`",
        "- Recommended migration tool: Flyway",
        "- First script: `create_foundation_tables`",
        "",
        "## Database Script Meaning",
        "",
        "A database script is a `.sql` file that creates or changes database objects such as database, tables, indexes, seed data, triggers and stored procedures.",
        "",
        "## Table Meaning",
        "",
        "A table is the actual structure in MySQL where data is stored. Example: `tenants`, `hospitals`, `branches`, `patients`.",
        "",
        "## Database Creation SQL",
        "",
        "```sql",
        "CREATE DATABASE IF NOT EXISTS plasmit_hms",
        "  CHARACTER SET utf8mb4",
        "  COLLATE utf8mb4_0900_ai_ci;",
        "",
        "USE plasmit_hms;",
        "```",
        "",
        "## High-Level Script Order",
        "",
    ]
    lines.extend(md_table(["Order", "Script Name", "Purpose", "Main Objects", "Note"], SCRIPT_ORDER))
    lines.extend(["", "## First Tables to Create", ""])
    lines.extend(md_table(["Order", "Table", "Purpose", "Depends On", "Why Important"], FIRST_TABLES))
    lines.extend(["", "## Module-Wise Creation Flow", ""])
    lines.extend(md_table(["Module", "Table Flow", "Business Meaning"], MODULE_ORDER))
    lines.extend([
        "",
        "## Flyway Naming Convention",
        "",
        "```text",
        "V202606030001__create_foundation_tables.sql",
        "V202606030002__create_user_role_permission_tables.sql",
        "V202606030003__create_patient_tables.sql",
        "V202606030004__create_appointment_tables.sql",
        "V202606030005__create_encounter_clinical_tables.sql",
        "V202606030006__create_lab_radiology_tables.sql",
        "V202606030007__create_pharmacy_tables.sql",
        "V202606030008__create_billing_tables.sql",
        "V202606030009__create_fhir_hl7_tables.sql",
        "R__seed_default_master_data.sql",
        "```",
        "",
        "## Final Recommendation",
        "",
        "Create `plasmit_hms` first. Then create `tenants` first. After that create `hospitals`, `branches`, `departments`, `healthcare_services`, `users`, `roles`, `permissions`, `role_permissions`, and `user_branch_access`.",
    ])
    return "\n".join(lines)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(DOCX)
    MD.write_text(build_markdown(), encoding="utf-8")
    with ZipFile(DOCX) as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError(f"Bad docx part: {bad}")
    print(f"created: {DOCX}")
    print(f"created: {MD}")


if __name__ == "__main__":
    main()
