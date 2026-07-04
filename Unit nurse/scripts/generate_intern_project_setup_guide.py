from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DOCX = DOCS / "Plasmit-HMS-Frontend-Intern-Setup-Guide.docx"
MD = DOCS / "plasmit-hms-frontend-intern-setup-guide.md"


REPO_URL = "https://github.com/Aman305-bit/plasmit-web-application-UI.git"
PROJECT_FOLDER = "plasmit-web-application-UI"


def rgb(hex_color: str) -> tuple[int, int, int]:
    hex_color = hex_color.replace("#", "")
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border(cell, color: str = "CBD5E1") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Aptos Display"
        run.font.bold = True
        run.font.color.rgb = RGBColor(*rgb("14325C" if level <= 2 else "2563EB"))


def para(doc: Document, text: str, bold: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10.3)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*rgb("334155"))


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(item)
        run.font.name = "Aptos"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(*rgb("334155"))


def code(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, "0F172A")
    border(cell, "1E293B")
    paragraph = cell.paragraphs[0]
    for line in text.strip().splitlines():
        run = paragraph.add_run(line + "\n")
        run.font.name = "Consolas"
        run.font.size = Pt(8.4)
        run.font.color.rgb = RGBColor(*rgb("E2E8F0"))
    doc.add_paragraph()


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = tbl.cell(0, index)
        shade(cell, "14325C")
        border(cell)
        run = cell.paragraphs[0].add_run(header)
        run.font.name = "Aptos"
        run.font.size = Pt(8.2)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            border(cell)
            if row_index % 2 == 0:
                shade(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Aptos"
            run.font.size = Pt(8.0)
            run.font.color.rgb = RGBColor(*rgb("334155"))
    doc.add_paragraph()


def callout(doc: Document, title: str, items: list[str]) -> None:
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    shade(cell, "EAF4FF")
    border(cell, "BBD4F8")
    run = cell.paragraphs[0].add_run(title)
    run.font.name = "Aptos Display"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(*rgb("14325C"))
    for item in items:
        p = cell.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        r.font.name = "Aptos"
        r.font.size = Pt(9.4)
        r.font.color.rgb = RGBColor(*rgb("334155"))
    doc.add_paragraph()


def build_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    cover = doc.add_table(rows=1, cols=1)
    cell = cover.cell(0, 0)
    shade(cell, "14325C")
    border(cell, "14325C")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Plasmit HMS Frontend Project")
    run.font.name = "Aptos Display"
    run.font.size = Pt(21)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("New Intern Setup Guide")
    r2.font.name = "Aptos Display"
    r2.font.size = Pt(15)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(*rgb("D8E8FF"))
    doc.add_paragraph()

    table(doc, ["Project", "Type", "Framework", "Repository"], [[
        "Plasmit Hospital HMS Frontend",
        "Frontend web application",
        "Next.js + TypeScript",
        REPO_URL,
    ]])

    heading(doc, "1. Purpose", 1)
    para(doc, "This guide explains how a new intern can install the required tools, download the project from Git, open it in VS Code, install dependencies, and run the Next.js frontend project locally.")

    heading(doc, "2. Required Software", 1)
    table(doc, ["Software", "Purpose", "Download Link"], [
        ["Git", "Download project from GitHub and manage code changes", "https://git-scm.com/downloads"],
        ["Node.js LTS", "Run Next.js project and npm commands", "https://nodejs.org"],
        ["VS Code", "Code editor", "https://code.visualstudio.com"],
        ["Google Chrome", "Browser for testing", "https://www.google.com/chrome"],
    ])

    heading(doc, "3. Install Git", 1)
    bullets(doc, [
        "Download Git from https://git-scm.com/downloads.",
        "Run the installer.",
        "Keep default options unless your senior developer asks for changes.",
        "After installation, open Command Prompt or PowerShell.",
        "Check Git installation using the command below.",
    ])
    code(doc, "git --version")

    heading(doc, "4. Install Node.js", 1)
    bullets(doc, [
        "Download Node.js LTS from https://nodejs.org.",
        "Install Node.js with default options.",
        "After installation, open a new terminal.",
        "Check Node and npm versions.",
    ])
    code(doc, """
node -v
npm -v
""")

    heading(doc, "5. Install VS Code", 1)
    bullets(doc, [
        "Download VS Code from https://code.visualstudio.com.",
        "Install it with default options.",
        "Recommended extensions: ESLint, Prettier, Tailwind CSS IntelliSense, GitLens.",
    ])

    heading(doc, "6. Clone the Project from GitHub", 1)
    para(doc, "Choose a folder where you want to keep the project. Example: Desktop or Documents.")
    code(doc, f"""
cd Desktop
git clone {REPO_URL}
cd {PROJECT_FOLDER}
""")
    callout(doc, "If repository is private", [
        "Ask the project owner to add your GitHub account as collaborator.",
        "Login to GitHub when Git asks for authentication.",
        "Use your GitHub username and token if password login is not accepted.",
    ])

    heading(doc, "7. Open Project in VS Code", 1)
    code(doc, "code .")
    para(doc, "If the command does not work, open VS Code manually, then click File > Open Folder and select the project folder.")

    heading(doc, "8. Install Project Dependencies", 1)
    para(doc, "Run this command inside the project folder.")
    code(doc, "npm install")
    para(doc, "This will create the node_modules folder. Do not manually edit node_modules.")

    heading(doc, "9. Environment File", 1)
    bullets(doc, [
        "Check if the project has .env.example.",
        "Create a new file named .env.local if required.",
        "Copy values from .env.example and ask the senior developer for actual values.",
        "Never commit .env or .env.local to Git.",
    ])
    code(doc, """
copy .env.example .env.local
""")
    para(doc, "If there is no .env.example and the project runs without environment variables, skip this step.")

    heading(doc, "10. Run the Project", 1)
    code(doc, "npm run dev")
    para(doc, "After the server starts, open the browser and visit:")
    code(doc, "http://localhost:3000")

    heading(doc, "11. Useful Commands", 1)
    table(doc, ["Command", "Purpose"], [
        ["npm run dev", "Start local development server"],
        ["npm run build", "Create production build"],
        ["npm run typecheck", "Check TypeScript errors"],
        ["npm run lint", "Check linting/code style"],
        ["git status", "Check changed files"],
        ["git pull origin dev", "Get latest code from dev branch"],
        ["git checkout -b feature/my-task-name", "Create new feature branch"],
    ])

    heading(doc, "12. Daily Git Workflow for Intern", 1)
    code(doc, """
git checkout dev
git pull origin dev
git checkout -b feature/your-task-name

# after code changes
git status
git add .
git commit -m "feat: describe your change"
git push origin feature/your-task-name
""")
    bullets(doc, [
        "Do not directly push to dev unless senior developer allows it.",
        "Create a feature branch for every task.",
        "Commit only related files.",
        "Do not commit node_modules, .next, .env, logs, or temporary files.",
    ])

    heading(doc, "13. Project Folder Overview", 1)
    table(doc, ["Folder/File", "Meaning"], [
        ["src/app", "Next.js App Router pages and routes"],
        ["src/features", "Module-wise frontend code"],
        ["src/components", "Reusable UI components"],
        ["src/data", "Static app data such as navigation"],
        ["public", "Images and static assets"],
        ["docs", "Project documentation"],
        ["database", "Database SQL scripts and planning documents"],
        ["package.json", "Project scripts and dependencies"],
    ])

    heading(doc, "14. Common Problems and Fixes", 1)
    table(doc, ["Problem", "Solution"], [
        ["npm is not recognized", "Install Node.js LTS and reopen terminal."],
        ["git is not recognized", "Install Git and reopen terminal."],
        ["Port 3000 already in use", "Stop old server or run Next.js on another port."],
        ["Module not found", "Run npm install again."],
        ["TypeScript error", "Read terminal error and ask senior developer if unclear."],
        ["Private repo access denied", "Ask owner to add your GitHub account."],
    ])

    heading(doc, "15. Final Checklist", 1)
    bullets(doc, [
        "Git installed.",
        "Node.js LTS installed.",
        "VS Code installed.",
        "Project cloned from GitHub.",
        "Project opened in VS Code.",
        "npm install completed.",
        "npm run dev started successfully.",
        "Website opened at http://localhost:3000.",
        "Intern understands git status, pull, branch, commit, and push basics.",
    ])

    return doc


def build_markdown() -> str:
    return f"""# Plasmit HMS Frontend Project - New Intern Setup Guide

## 1. Purpose

This guide explains how a new intern can install tools, download the project from GitHub, open it in VS Code, install dependencies, and run the Next.js frontend locally.

## 2. Required Software

| Software | Purpose | Link |
| --- | --- | --- |
| Git | Clone and manage code | https://git-scm.com/downloads |
| Node.js LTS | Run Next.js and npm | https://nodejs.org |
| VS Code | Code editor | https://code.visualstudio.com |
| Google Chrome | Browser testing | https://www.google.com/chrome |

## 3. Install Git

```bash
git --version
```

## 4. Install Node.js

```bash
node -v
npm -v
```

## 5. Clone the Project

```bash
cd Desktop
git clone {REPO_URL}
cd {PROJECT_FOLDER}
```

If the repository is private, ask the project owner to add your GitHub account.

## 6. Open in VS Code

```bash
code .
```

If this command does not work, open VS Code manually and use File > Open Folder.

## 7. Install Dependencies

```bash
npm install
```

## 8. Environment File

If `.env.example` exists:

```bash
copy .env.example .env.local
```

Ask the senior developer for actual environment values. Never commit `.env` or `.env.local`.

## 9. Run the Project

```bash
npm run dev
```

Open:

```text
http://localhost:3000
```

## 10. Useful Commands

| Command | Purpose |
| --- | --- |
| npm run dev | Start local development server |
| npm run build | Production build |
| npm run typecheck | Check TypeScript errors |
| npm run lint | Check linting |
| git status | Check changed files |
| git pull origin dev | Get latest dev code |
| git checkout -b feature/my-task-name | Create feature branch |

## 11. Daily Git Workflow

```bash
git checkout dev
git pull origin dev
git checkout -b feature/your-task-name

# after code changes
git status
git add .
git commit -m "feat: describe your change"
git push origin feature/your-task-name
```

## 12. Project Folder Overview

| Folder/File | Meaning |
| --- | --- |
| src/app | Next.js App Router pages/routes |
| src/features | Module-wise frontend code |
| src/components | Reusable UI components |
| src/data | Static app data |
| public | Static assets |
| docs | Documentation |
| database | SQL/database scripts |
| package.json | Scripts and dependencies |

## 13. Common Problems

| Problem | Solution |
| --- | --- |
| npm is not recognized | Install Node.js LTS and reopen terminal |
| git is not recognized | Install Git and reopen terminal |
| Port 3000 already in use | Stop old server or run on another port |
| Module not found | Run npm install again |
| Private repo access denied | Ask owner to add your GitHub account |

## 14. Final Checklist

- Git installed
- Node.js LTS installed
- VS Code installed
- Project cloned
- Project opened in VS Code
- `npm install` completed
- `npm run dev` started
- Website opened at `http://localhost:3000`
"""


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    doc = build_doc()
    doc.save(DOCX)
    MD.write_text(build_markdown(), encoding="utf-8")
    with ZipFile(DOCX) as zf:
        bad = zf.testzip()
        if bad:
            raise RuntimeError(f"Bad docx part: {bad}")
    print(f"created: {DOCX}")
    print(f"created: {MD}")


if __name__ == "__main__":
    main()
