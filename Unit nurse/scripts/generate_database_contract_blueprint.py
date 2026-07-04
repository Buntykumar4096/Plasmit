from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets" / "database-contract"
DOCX = DOCS / "Plasmit-HMS-High-Level-Database-Contract-Table-Relationship-FHIR-Mapping.docx"
MD = DOCS / "plasmit-hms-high-level-database-contract-fhir-mapping.md"


def rgb(hex_color: str) -> tuple[int, int, int]:
    hex_color = hex_color.replace("#", "")
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def fill(value: str | tuple[int, int, int]) -> str | tuple[int, int, int]:
    if isinstance(value, tuple):
        return value
    if len(value.replace("#", "")) == 6:
        return rgb(value)
    return value


def font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    name = "arialbd.ttf" if bold else "arial.ttf"
    for candidate in [Path("C:/Windows/Fonts") / name, Path("C:/Windows/Fonts/calibri.ttf")]:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def draw_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    box: tuple[int, int, int, int],
    size: int = 22,
    color: str = "334155",
    bold: bool = False,
    align: str = "center",
) -> None:
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 24
    f = font(size, bold)
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        trial = f"{line} {word}".strip()
        bbox = draw.textbbox((0, 0), trial, font=f)
        if bbox[2] - bbox[0] <= max_width:
            line = trial
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    heights = [draw.textbbox((0, 0), item, font=f)[3] for item in lines]
    total_height = sum(heights) + max(0, len(lines) - 1) * 7
    y = y1 + max(0, (y2 - y1 - total_height) // 2)
    for item, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), item, font=f)
        width = bbox[2] - bbox[0]
        x = x1 + 16 if align == "left" else x1 + (x2 - x1 - width) // 2
        draw.text((x, y), item, font=f, fill=fill(color))
        y += height + 7


def box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    body: str = "",
    bg: str = "FFFFFF",
    outline: str = "CBD5E1",
    title_color: str = "14325C",
) -> None:
    draw.rounded_rectangle(xy, radius=20, fill=fill(bg), outline=fill(outline), width=3)
    x1, y1, x2, y2 = xy
    draw_text(draw, title, (x1 + 8, y1 + 10, x2 - 8, y1 + 66), size=24, color=title_color, bold=True)
    if body:
        draw_text(draw, body, (x1 + 12, y1 + 70, x2 - 12, y2 - 10), size=19, color="475569")


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str = "2563EB") -> None:
    draw.line([start, end], fill=fill(color), width=5)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex >= sx else -1
        points = [(ex, ey), (ex - 18 * direction, ey - 10), (ex - 18 * direction, ey + 10)]
    else:
        direction = 1 if ey >= sy else -1
        points = [(ex, ey), (ex - 10, ey - 18 * direction), (ex + 10, ey - 18 * direction)]
    draw.polygon(points, fill=fill(color))


def base(title: str, subtitle: str = "") -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (1600, 900), rgb("F8FAFC"))
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 1600, 118), fill=rgb("14325C"))
    draw.text((48, 25), title, font=font(40, True), fill=rgb("FFFFFF"))
    if subtitle:
        draw.text((50, 80), subtitle, font=font(22), fill=rgb("D8E8FF"))
    return img, draw


def save_img(img: Image.Image, name: str) -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    path = ASSETS / name
    img.save(path, "PNG", quality=95)
    return path


def create_diagrams() -> dict[str, Path]:
    diagrams: dict[str, Path] = {}

    img, draw = base("Multi-Tenant Database Contract Hierarchy", "Tenant, hospital, branch, patient and encounter are the core contract anchors.")
    nodes = [
        ("Tenant / Hospital Group", "subscription and ownership", 690, 160),
        ("Hospital", "legal healthcare entity", 690, 275),
        ("Branch / Facility / Location", "care delivery location", 690, 390),
        ("Department / Service", "OPD, IPD, ER, Lab, Pharmacy", 270, 520),
        ("Users / Doctors / Staff", "RBAC and branch access", 690, 520),
        ("Patient", "hospital-level identity", 1110, 520),
        ("Encounter / Visit", "branch-level care event", 1110, 650),
    ]
    for title, body, cx, cy in nodes:
        box(draw, (cx - 190, cy - 45, cx + 190, cy + 45), title, body)
    arrow(draw, (690, 205), (690, 230), color="0F766E")
    arrow(draw, (690, 320), (690, 345), color="0F766E")
    arrow(draw, (570, 430), (340, 475), color="64748B")
    arrow(draw, (690, 435), (690, 475), color="64748B")
    arrow(draw, (810, 430), (1045, 475), color="64748B")
    arrow(draw, (1110, 565), (1110, 600), color="2563EB")
    draw_text(draw, "Contract rule: patient can visit many branches, but encounter, orders, billing and stock are branch-level records.", (110, 770, 1490, 840), size=29, color="B91C1C", bold=True)
    diagrams["hierarchy"] = save_img(img, "01_multitenant_hierarchy.png")

    img, draw = base("Foundation Table Relationship Contract", "Foundation tables must be created before patient, clinical, billing and integration modules.")
    layout = [
        ("tenants", "PK id", 140, 230),
        ("hospitals", "FK tenant_id", 470, 230),
        ("branches", "FK hospital_id", 800, 230),
        ("departments", "FK hospital_id / branch_id", 1130, 230),
        ("healthcare_services", "FK department_id", 1130, 420),
        ("users", "FK tenant_id, hospital_id", 470, 520),
        ("roles", "RBAC role", 140, 520),
        ("permissions", "RBAC permission", 140, 710),
        ("role_permissions", "role_id + permission_id", 470, 710),
        ("user_branch_access", "user_id + branch_id", 800, 520),
    ]
    for title, body, x, y in layout:
        box(draw, (x, y, x + 250, y + 95), title, body, bg="FFFFFF")
    arrow(draw, (390, 278), (465, 278))
    arrow(draw, (720, 278), (795, 278))
    arrow(draw, (1050, 278), (1125, 278))
    arrow(draw, (1255, 325), (1255, 415), color="0F766E")
    arrow(draw, (595, 325), (595, 515), color="64748B")
    arrow(draw, (720, 565), (795, 565), color="64748B")
    arrow(draw, (265, 615), (265, 705), color="64748B")
    arrow(draw, (390, 755), (465, 755), color="64748B")
    diagrams["foundation"] = save_img(img, "02_foundation_relationship_contract.png")

    img, draw = base("Encounter-Centric Clinical Contract", "Encounter connects clinical, lab, radiology, prescription, billing and documents.")
    box(draw, (650, 180, 950, 295), "encounters", "PK id\nFK patient_id, branch_id", bg="EAF4FF")
    satellites = [
        ("patients", "hospital-level", 160, 190),
        ("appointments", "may create encounter", 160, 420),
        ("vitals", "Observation", 520, 420),
        ("diagnoses", "Condition", 810, 420),
        ("clinical_notes", "DocumentReference", 1100, 420),
        ("lab_orders", "ServiceRequest", 160, 650),
        ("prescriptions", "MedicationRequest", 520, 650),
        ("billing_invoices", "Account / ChargeItem", 810, 650),
        ("documents", "DocumentReference", 1100, 650),
    ]
    for title, body, x, y in satellites:
        box(draw, (x, y, x + 250, y + 100), title, body)
    for x, y in [(410, 240), (410, 470), (650, 470), (810, 420), (1100, 470), (410, 690), (650, 690), (810, 650), (1100, 690)]:
        target = (800, 300) if y < 550 else (800, 295)
        arrow(draw, (x, y), target, color="2563EB")
    diagrams["encounter"] = save_img(img, "03_encounter_centric_clinical_contract.png")

    img, draw = base("Lab Order to Diagnostic Report Contract", "Lab workflow must preserve order, item, sample, result and report relationships.")
    steps = [
        ("Encounter", "source visit"),
        ("Lab Order", "FK encounter_id"),
        ("Order Items", "test-wise rows"),
        ("Test Master", "parameter template"),
        ("Sample / Specimen", "collection lifecycle"),
        ("Result", "verified values"),
        ("Result Parameters", "Observation rows"),
        ("Diagnostic Report", "final report"),
    ]
    x = 55
    for index, (title, body) in enumerate(steps):
        y = 280 if index < 4 else 550
        if index == 4:
            x = 200
        box(draw, (x, y, x + 210, y + 110), title, body, bg="FFFFFF" if index != 6 else "E8F7F3")
        if index < 3:
            arrow(draw, (x + 214, y + 55), (x + 252, y + 55))
        if index == 3:
            arrow(draw, (x + 105, y + 115), (x - 540, 545), color="64748B")
        if index > 4 and index < 7:
            arrow(draw, (x + 214, y + 55), (x + 252, y + 55))
        x += 260
    diagrams["lab"] = save_img(img, "04_lab_order_result_contract.png")

    img, draw = base("Pharmacy and Inventory Contract", "Prescription, stock, sale, dispense and stock movement must stay connected.")
    items = [
        ("Medicine Master", "medicine catalog", 90, 220),
        ("Medicine Batches", "expiry, batch, MRP", 430, 220),
        ("Branch Stock", "branch_id required", 770, 220),
        ("Prescription", "MedicationRequest", 90, 560),
        ("Pharmacy Sale", "invoice/receipt", 430, 560),
        ("Sale Items", "batch-wise issue", 770, 560),
        ("Stock Movements", "every change tracked", 1110, 390),
    ]
    for title, body, x, y in items:
        box(draw, (x, y, x + 270, y + 115), title, body)
    arrow(draw, (360, 278), (425, 278))
    arrow(draw, (700, 278), (765, 278))
    arrow(draw, (360, 618), (425, 618), color="0F766E")
    arrow(draw, (700, 618), (765, 618), color="0F766E")
    arrow(draw, (1040, 278), (1105, 445), color="64748B")
    arrow(draw, (1040, 618), (1105, 500), color="64748B")
    diagrams["pharmacy"] = save_img(img, "05_pharmacy_inventory_contract.png")

    img, draw = base("Billing, Payment and Insurance Contract", "Every charge, payment, refund and claim must be traceable to patient and encounter.")
    items = [
        ("Patient", "FK patient_id", 95, 230),
        ("Encounter", "FK encounter_id", 95, 500),
        ("Billing Invoice", "branch-level invoice", 470, 360),
        ("Invoice Items", "ChargeItem rows", 850, 230),
        ("Payments", "receipt events", 850, 500),
        ("Payment Allocations", "invoice-payment link", 1180, 500),
        ("Insurance Claims", "payer claim", 1180, 230),
    ]
    for title, body, x, y in items:
        box(draw, (x, y, x + 260, y + 112), title, body)
    arrow(draw, (360, 286), (465, 405))
    arrow(draw, (360, 555), (465, 430))
    arrow(draw, (735, 405), (845, 286))
    arrow(draw, (735, 430), (845, 555))
    arrow(draw, (1110, 555), (1175, 555))
    arrow(draw, (1110, 286), (1175, 286))
    diagrams["billing"] = save_img(img, "06_billing_payment_insurance_contract.png")

    img, draw = base("FHIR and HL7 Integration Readiness Contract", "FHIR and HL7 mappings are contract-managed integration layers.")
    boxes = [
        ("Internal Normalized Tables", "patients, encounters, orders, results, billing", 90, 260),
        ("FHIR Mapping", "internal record -> FHIR resource", 500, 170),
        ("FHIR API / Store", "FHIR JSON, audit, sync status", 900, 170),
        ("HL7 Message Log", "ADT, ORM, ORU, SIU, DFT, MDM", 500, 520),
        ("External Systems", "EHR, LIS, RIS, Insurance, Devices", 900, 520),
    ]
    for title, body, x, y in boxes:
        box(draw, (x, y, x + 330, y + 130), title, body, bg="EAF4FF" if "FHIR" in title else "FFFFFF")
    arrow(draw, (425, 325), (495, 235))
    arrow(draw, (835, 235), (895, 235))
    arrow(draw, (425, 325), (495, 585), color="0F766E")
    arrow(draw, (835, 585), (895, 585), color="0F766E")
    draw_text(draw, "Contract rule: module developers own internal data; integration team owns FHIR/HL7 resource mapping and external identifiers.", (120, 740, 1480, 820), size=28, color="14325C", bold=True)
    diagrams["integration"] = save_img(img, "07_fhir_hl7_integration_contract.png")

    img, draw = base("Transactional to Reporting Contract", "Summary tables are generated outputs, never source of truth.")
    box(draw, (100, 310, 420, 510), "Normalized Transaction Tables", "patient, encounter, lab, pharmacy, billing", bg="EAF4FF")
    box(draw, (570, 310, 890, 510), "Aggregation Jobs", "scheduled or event-driven refresh", bg="FFF7E6")
    box(draw, (1040, 190, 1400, 390), "Denormalized Summary Tables", "daily revenue, visit count, bed occupancy", bg="E8F7F3")
    box(draw, (1040, 520, 1400, 720), "Dashboards / Reports / AI", "fast analytics and command center", bg="FFFFFF")
    arrow(draw, (425, 410), (565, 410))
    arrow(draw, (895, 390), (1035, 300), color="0F766E")
    arrow(draw, (895, 430), (1035, 620), color="0F766E")
    diagrams["reporting"] = save_img(img, "08_reporting_contract.png")

    return diagrams


def shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def border(cell, hex_color: str = "CBD5E1") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), hex_color)


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Plasmit HMS | High-Level Database Contract, Table Relationship Blueprint and FHIR Mapping | Confidential")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(*rgb("64748B"))
    return doc


def heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Aptos Display"
        run.font.bold = True
        run.font.color.rgb = RGBColor(*rgb("14325C" if level <= 2 else "2563EB"))


def para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Aptos"
    run.font.size = Pt(10.2)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*rgb("334155"))


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        run.font.name = "Aptos"
        run.font.size = Pt(9.4)
        run.font.color.rgb = RGBColor(*rgb("334155"))


def callout(doc: Document, title: str, items: list[str], bg: str = "EAF4FF", title_color: str = "14325C") -> None:
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    shade(cell, bg)
    border(cell, "BBD4F8")
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.font.name = "Aptos Display"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(*rgb(title_color))
    for item in items:
        bp = cell.add_paragraph(style="List Bullet")
        br = bp.add_run(item)
        br.font.name = "Aptos"
        br.font.size = Pt(9.2)
        br.font.color.rgb = RGBColor(*rgb("334155"))
    doc.add_paragraph()


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = tbl.cell(0, idx)
        shade(cell, "14325C")
        border(cell)
        r = cell.paragraphs[0].add_run(header)
        r.font.name = "Aptos"
        r.font.size = Pt(7.8)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            border(cell)
            if row_index % 2 == 0:
                shade(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            r.font.name = "Aptos"
            r.font.size = Pt(7.1 if len(headers) >= 7 else 8.0)
            r.font.color.rgb = RGBColor(*rgb("334155"))
    doc.add_paragraph()


def picture(doc: Document, path: Path, caption: str) -> None:
    doc.add_picture(str(path), width=Inches(7.1))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    r.font.name = "Aptos"
    r.font.size = Pt(8.1)
    r.font.italic = True
    r.font.color.rgb = RGBColor(*rgb("64748B"))


def code(doc: Document, text: str) -> None:
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.cell(0, 0)
    shade(cell, "0F172A")
    border(cell, "1E293B")
    p = cell.paragraphs[0]
    for line in text.strip().splitlines():
        r = p.add_run(line + "\n")
        r.font.name = "Consolas"
        r.font.size = Pt(8.0)
        r.font.color.rgb = RGBColor(*rgb("E2E8F0"))
    doc.add_paragraph()


FOUNDATION_RELATIONSHIPS = [
    ["tenants", "hospitals", "1:N", "hospitals.tenant_id", "One tenant can own many hospitals.", "Organization"],
    ["hospitals", "branches", "1:N", "branches.hospital_id", "One hospital can have many branches/facilities.", "Location"],
    ["hospitals", "departments", "1:N", "departments.hospital_id", "Hospital defines departments.", "HealthcareService"],
    ["branches", "departments", "1:N optional", "departments.branch_id", "Branch-specific departments when required.", "HealthcareService"],
    ["departments", "healthcare_services", "1:N", "healthcare_services.department_id", "Department provides services.", "HealthcareService"],
    ["users", "user_branch_access", "1:N", "user_branch_access.user_id", "User can access multiple authorized branches.", "PractitionerRole"],
    ["branches", "user_branch_access", "1:N", "user_branch_access.branch_id", "Branch access is explicitly assigned.", "Location"],
    ["roles", "role_permissions", "1:N", "role_permissions.role_id", "Role receives permissions.", "N/A"],
    ["permissions", "role_permissions", "1:N", "role_permissions.permission_id", "Permission can be part of many roles.", "N/A"],
]


MODULE_TABLES = [
    ["Foundation", "tenants, hospitals, branches, departments, healthcare_services, users, roles, permissions, role_permissions, user_branch_access", "Organization, Location, HealthcareService, Practitioner, PractitionerRole"],
    ["Patient", "patients, patient_identifiers, patient_addresses, patient_contacts, patient_allergies, patient_insurance, patient_documents, patient_consent_records, patient_record_access_logs", "Patient, RelatedPerson, AllergyIntolerance, DocumentReference, Consent, AuditEvent"],
    ["Appointment", "appointment_slots, appointments, appointment_status_history, appointment_reminders, appointment_cancellations", "Appointment, Patient, Practitioner, Location, AuditEvent"],
    ["Encounter / Clinical", "encounters, encounter_participants, encounter_status_history, vitals, diagnoses, clinical_notes, clinical_documents, procedures, care_plans", "Encounter, Observation, Condition, DocumentReference, Procedure, CarePlan"],
    ["Prescription / Medication", "prescriptions, prescription_items, medicine_master, medication_instructions, medication_administration_records, pharmacy_dispense, pharmacy_dispense_items", "MedicationRequest, Medication, MedicationDispense, MedicationAdministration"],
    ["Lab / Diagnostics", "lab_test_master, lab_test_parameters, lab_orders, lab_order_items, lab_samples, lab_results, lab_result_parameters, diagnostic_reports, lab_result_review_history", "ServiceRequest, Observation, DiagnosticReport, Specimen"],
    ["Radiology", "radiology_test_master, radiology_orders, radiology_order_items, radiology_reports, imaging_studies, radiology_report_review_history", "ServiceRequest, DiagnosticReport, ImagingStudy, DocumentReference"],
    ["Pharmacy / Inventory", "medicine_categories, medicine_batches, pharmacy_stock, stock_movements, pharmacy_sales, pharmacy_sale_items, pharmacy_returns, pharmacy_return_items, purchase_orders, goods_receipts", "Medication, MedicationDispense"],
    ["Billing / Payment / Insurance", "billing_invoices, billing_invoice_items, payments, payment_allocations, refunds, insurance_policies, insurance_claims, insurance_claim_items, claim_documents, claim_status_history", "Account, ChargeItem, Claim, ExplanationOfBenefit, Coverage"],
    ["IPD / Ward / Bed", "admissions, wards, rooms, beds, bed_allocations, nursing_notes, nursing_tasks, discharge_plans, discharge_summaries", "Encounter, Location, DocumentReference, Composition"],
    ["Surgery / OT", "operation_theatres, surgery_cases, surgery_team_members, surgery_checklists, anesthesia_records, surgery_notes, post_operation_notes", "Procedure, Location, Observation, DocumentReference"],
    ["Documents", "documents, document_versions, document_access_logs, document_signatures", "DocumentReference, Composition"],
    ["Audit / Compliance", "audit_logs, user_login_history, patient_record_access_logs, consent_records, break_glass_access_logs, data_export_logs", "AuditEvent, Consent, Provenance"],
    ["FHIR Integration", "fhir_resource_mapping, fhir_resource_store, fhir_api_audit_log, fhir_sync_status, fhir_profile_registry, fhir_code_system_mapping, fhir_value_set_mapping", "All mapped FHIR resources"],
    ["HL7 Integration", "hl7_message_log, hl7_message_error_log, hl7_external_identifier_mapping, hl7_patient_mapping, hl7_order_mapping, hl7_result_mapping", "ADT, ORM, ORU, SIU, DFT, MDM message readiness"],
    ["Reporting", "daily_branch_revenue_summary, daily_patient_visit_summary, doctor_performance_summary, department_collection_summary, lab_test_volume_summary, pharmacy_stock_snapshot, bed_occupancy_summary, patient_journey_summary", "N/A"],
]


RELATIONSHIP_MATRIX = [
    ["Foundation", "tenants", "id", "None", "None", "hospitals", "1:N", "Global/Tenant", "No", "Organization", "Foundation Team"],
    ["Foundation", "hospitals", "id", "tenant_id", "tenants", "branches, departments, patients", "1:N", "Tenant", "No", "Organization", "Foundation Team"],
    ["Foundation", "branches", "id", "tenant_id, hospital_id", "hospitals", "appointments, encounters, pharmacy_stock", "1:N", "Hospital", "Yes", "Location", "Foundation Team"],
    ["Foundation", "users", "id", "tenant_id, hospital_id", "hospitals", "appointments, encounters, audit_logs", "1:N", "Hospital", "Optional", "Practitioner", "Foundation Team"],
    ["Patient", "patients", "id", "tenant_id, hospital_id", "hospitals", "encounters, appointments, documents", "1:N", "Hospital", "No by default", "Patient", "Patient Team"],
    ["Patient", "patient_identifiers", "id", "patient_id", "patients", "None", "N:1", "Hospital", "No", "Patient.identifier", "Patient Team"],
    ["Appointment", "appointments", "id", "patient_id, branch_id, doctor_id", "patients, branches, users", "encounters, reminders", "N:1 / 1:N", "Hospital", "Yes", "Appointment", "Appointment Team"],
    ["Encounter", "encounters", "id", "patient_id, branch_id, appointment_id", "patients, branches, appointments", "vitals, diagnoses, lab_orders, prescriptions, invoices", "1:N", "Hospital", "Yes", "Encounter", "Clinical Team"],
    ["Clinical", "vitals", "id", "patient_id, encounter_id", "patients, encounters", "None", "N:1", "Hospital", "Yes", "Observation", "Clinical Team"],
    ["Clinical", "diagnoses", "id", "patient_id, encounter_id", "patients, encounters", "None", "N:1", "Hospital", "Yes", "Condition", "Clinical Team"],
    ["Clinical", "clinical_notes", "id", "patient_id, encounter_id, doctor_id", "patients, encounters, users", "documents", "N:1", "Hospital", "Yes", "DocumentReference", "Clinical Team"],
    ["Prescription", "prescriptions", "id", "patient_id, encounter_id, doctor_id", "patients, encounters, users", "prescription_items, pharmacy_dispense", "1:N", "Hospital", "Yes", "MedicationRequest", "Clinical Team"],
    ["Prescription", "prescription_items", "id", "prescription_id, medicine_id", "prescriptions, medicine_master", "pharmacy_dispense_items", "N:1", "Hospital", "Yes", "MedicationRequest.dosage", "Clinical/Pharmacy Team"],
    ["Lab", "lab_orders", "id", "patient_id, encounter_id, branch_id", "patients, encounters, branches", "lab_order_items", "1:N", "Hospital", "Yes", "ServiceRequest", "Lab Team"],
    ["Lab", "lab_order_items", "id", "lab_order_id, lab_test_id", "lab_orders, lab_test_master", "lab_samples, lab_results", "1:N", "Hospital", "Yes", "ServiceRequest", "Lab Team"],
    ["Lab", "lab_results", "id", "lab_order_item_id, patient_id, encounter_id", "lab_order_items, patients, encounters", "lab_result_parameters, diagnostic_reports", "1:N", "Hospital", "Yes", "Observation", "Lab Team"],
    ["Radiology", "radiology_orders", "id", "patient_id, encounter_id, branch_id", "patients, encounters, branches", "radiology_order_items", "1:N", "Hospital", "Yes", "ServiceRequest", "Radiology Team"],
    ["Radiology", "radiology_reports", "id", "radiology_order_item_id", "radiology_order_items", "imaging_studies, documents", "1:N", "Hospital", "Yes", "DiagnosticReport", "Radiology Team"],
    ["Pharmacy", "medicine_master", "id", "tenant_id, hospital_id", "hospitals", "medicine_batches", "1:N", "Hospital", "No", "Medication", "Pharmacy Team"],
    ["Pharmacy", "medicine_batches", "id", "medicine_id", "medicine_master", "pharmacy_stock, sale_items", "1:N", "Hospital", "Optional", "Medication.batch", "Pharmacy Team"],
    ["Pharmacy", "pharmacy_stock", "id", "branch_id, medicine_batch_id", "branches, medicine_batches", "stock_movements", "1:N", "Hospital", "Yes", "Internal", "Pharmacy Team"],
    ["Pharmacy", "pharmacy_sales", "id", "patient_id, branch_id, prescription_id", "patients, branches, prescriptions", "pharmacy_sale_items", "1:N", "Hospital", "Yes", "MedicationDispense", "Pharmacy Team"],
    ["Billing", "billing_invoices", "id", "patient_id, encounter_id, branch_id", "patients, encounters, branches", "invoice_items, payments, claims", "1:N", "Hospital", "Yes", "Account", "Billing Team"],
    ["Billing", "billing_invoice_items", "id", "invoice_id", "billing_invoices", "None", "N:1", "Hospital", "Yes", "ChargeItem", "Billing Team"],
    ["Billing", "payments", "id", "tenant_id, hospital_id, branch_id", "billing_invoices via allocations", "payment_allocations", "N:M", "Hospital", "Yes", "PaymentReconciliation", "Billing Team"],
    ["Insurance", "insurance_claims", "id", "invoice_id, patient_id", "billing_invoices, patients", "claim_items, documents", "1:N", "Hospital", "Yes", "Claim", "Billing/Insurance Team"],
    ["IPD", "admissions", "id", "encounter_id, patient_id, branch_id", "encounters, patients, branches", "bed_allocations, nursing_notes, discharge_summaries", "1:N", "Hospital", "Yes", "Encounter", "IPD Team"],
    ["IPD", "beds", "id", "room_id, branch_id", "rooms, branches", "bed_allocations", "1:N", "Hospital", "Yes", "Location", "IPD Team"],
    ["Surgery", "surgery_cases", "id", "encounter_id, operation_theatre_id", "encounters, operation_theatres", "surgery_team_members, notes", "1:N", "Hospital", "Yes", "Procedure", "Surgery Team"],
    ["Documents", "documents", "id", "patient_id, encounter_id", "patients, encounters", "document_versions, access_logs", "1:N", "Hospital", "Optional", "DocumentReference", "Document Team"],
    ["Audit", "audit_logs", "id", "tenant_id, hospital_id, branch_id", "users, all entities", "None", "N:1", "Tenant/Hospital", "Optional", "AuditEvent", "Audit Team"],
    ["FHIR", "fhir_resource_mapping", "id", "internal_table_name, internal_record_id", "All mapped tables", "fhir_resource_store", "1:1 / 1:N", "Tenant/Hospital", "Optional", "All resources", "FHIR Team"],
    ["HL7", "hl7_message_log", "id", "tenant_id, hospital_id, branch_id", "External systems", "hl7 mappings", "1:N", "Tenant/Hospital", "Optional", "HL7 v2", "HL7 Team"],
    ["Reporting", "daily_branch_revenue_summary", "id", "tenant_id, hospital_id, branch_id", "billing_invoices, payments", "Dashboards", "Generated", "Hospital", "Yes", "N/A", "Reporting Team"],
]


FHIR_MAPPING = [
    ["Tenant / Hospital", "tenants, hospitals", "Organization", "Represent organization, tenant group or hospital.", "FHIR Team + Foundation"],
    ["Branch / Facility", "branches", "Location", "Represent facility, branch, ward, room or bed.", "FHIR Team + Foundation"],
    ["Department / Service", "departments, healthcare_services", "HealthcareService", "Expose healthcare service catalog.", "FHIR Team + Foundation"],
    ["Doctor / Staff", "users, practitioner_roles", "Practitioner, PractitionerRole", "Expose clinician and role context.", "FHIR Team + Security"],
    ["Patient", "patients, patient_identifiers, patient_addresses, patient_contacts", "Patient", "Expose patient demographics and identifiers.", "FHIR Team + Patient"],
    ["Consent", "patient_consent_records, consent_records", "Consent", "Consent tracking and data sharing permission.", "FHIR Team + Compliance"],
    ["Appointment", "appointments", "Appointment", "Appointment scheduling and participants.", "FHIR Team + Appointment"],
    ["Encounter", "encounters", "Encounter", "OPD, IPD, ER, day-care or virtual visit.", "FHIR Team + Clinical"],
    ["Vitals / Lab Values", "vitals, lab_results, lab_result_parameters", "Observation", "Clinical measurements and lab values.", "FHIR Team + Clinical/Lab"],
    ["Diagnosis", "diagnoses", "Condition", "Diagnosis/problem list.", "FHIR Team + Clinical"],
    ["Clinical Documents", "clinical_notes, documents", "DocumentReference, Composition", "Clinical notes, summaries and signed documents.", "FHIR Team + Documents"],
    ["Procedure / Surgery", "procedures, surgery_cases", "Procedure", "Procedure and surgery representation.", "FHIR Team + Surgery"],
    ["Care Plan", "care_plans", "CarePlan", "Patient care planning.", "FHIR Team + Clinical"],
    ["Prescription", "prescriptions, prescription_items", "MedicationRequest", "Medication orders.", "FHIR Team + Clinical/Pharmacy"],
    ["Medicine", "medicine_master", "Medication", "Medicine catalog/reference.", "FHIR Team + Pharmacy"],
    ["Dispensing", "pharmacy_dispense, pharmacy_sales", "MedicationDispense", "Dispensed medications.", "FHIR Team + Pharmacy"],
    ["Medication Admin", "medication_administration_records", "MedicationAdministration", "Inpatient medicine administration.", "FHIR Team + Nursing"],
    ["Lab/Radiology Order", "lab_orders, radiology_orders", "ServiceRequest", "Diagnostic, procedure and imaging orders.", "FHIR Team + Lab/Radiology"],
    ["Sample", "lab_samples", "Specimen", "Lab specimen lifecycle.", "FHIR Team + Lab"],
    ["Diagnostic Report", "diagnostic_reports, radiology_reports", "DiagnosticReport", "Lab/radiology report.", "FHIR Team + Diagnostics"],
    ["Imaging", "imaging_studies", "ImagingStudy", "Imaging study references.", "FHIR Team + Radiology"],
    ["Billing Account", "billing_invoices, patient_accounts", "Account", "Billing account context.", "FHIR Team + Billing"],
    ["Charges", "billing_invoice_items", "ChargeItem", "Charge line items.", "FHIR Team + Billing"],
    ["Insurance", "insurance_policies", "Coverage", "Insurance policy coverage.", "FHIR Team + Billing"],
    ["Claim", "insurance_claims", "Claim", "Insurance claim submission.", "FHIR Team + Billing"],
    ["Payment Explanation", "claim_status_history, payment_allocations", "ExplanationOfBenefit", "Claim/payment explanation when required.", "FHIR Team + Billing"],
    ["Audit", "audit_logs", "AuditEvent", "Security and data access audit.", "FHIR Team + Audit"],
]


OWNERSHIP = [
    ["Foundation", "Foundation Team", "tenants, hospitals, branches, departments, healthcare_services, users, roles, permissions", "None", "FHIR Team", "Database Architect"],
    ["Patient", "Patient Team", "patients, identifiers, addresses, contacts, allergies, documents, consent", "hospitals, branches, users", "FHIR Team", "Database Architect + FHIR Team"],
    ["Appointment", "Appointment Team", "appointment_slots, appointments, reminders, cancellations", "patients, branches, users", "FHIR Team", "Patient + Foundation"],
    ["Clinical / Encounter", "Clinical Team", "encounters, vitals, diagnoses, notes, procedures, care_plans", "patients, branches, users", "FHIR Team", "Patient + FHIR Team"],
    ["Lab / Radiology", "Diagnostics Team", "lab_orders, lab_results, diagnostic_reports, radiology_orders, imaging_studies", "encounters, billing", "FHIR/HL7 Team", "Clinical + Billing + Integration"],
    ["Pharmacy", "Pharmacy Team", "medicine_master, stock, sales, dispensing, returns", "prescriptions, billing", "FHIR Team", "Clinical + Billing"],
    ["Billing / Insurance", "Billing Team", "invoices, invoice_items, payments, refunds, policies, claims", "patients, encounters, services", "FHIR Team", "Database Architect + Finance"],
    ["IPD / Nursing", "IPD Team", "admissions, wards, rooms, beds, nursing_notes, discharge_summaries", "patients, encounters, billing", "FHIR Team", "Clinical + Billing"],
    ["Surgery / OT", "Surgery Team", "operation_theatres, surgery_cases, team, checklists, anesthesia, notes", "encounters, billing", "FHIR Team", "Clinical + Architect"],
    ["Documents", "Document Team", "documents, versions, access_logs, signatures", "patients, encounters, users", "FHIR Team", "Clinical + Security"],
    ["Audit / Compliance", "Audit Team", "audit_logs, login_history, access_logs, consent, break_glass, data_export", "all modules", "FHIR Team", "Compliance + Security"],
    ["FHIR Integration", "FHIR Team", "fhir_resource_mapping, store, audit, sync, profiles, code mapping", "all mapped modules", "FHIR Team", "Architect + Module Owners"],
    ["HL7 Integration", "HL7 Team", "hl7_message_log, errors, patient/order/result mappings", "patient, lab, radiology, billing", "Integration Team", "Architect + Module Owners"],
    ["Reporting", "Reporting Team", "summary and dashboard tables", "all modules", "N/A", "Architect + Module Owners"],
]


TENANT_SCOPE = [
    ["Global table", "id, code, name", "country_master, fhir_resource_type_master", "Readable by platform; no hospital isolation required."],
    ["Tenant-level table", "tenant_id", "subscription_plans, tenant_settings", "Tenant admin and platform scope only."],
    ["Hospital-level table", "tenant_id, hospital_id", "patients, departments, doctors, medicine_master", "Hospital cannot see another hospital's data."],
    ["Branch-level table", "tenant_id, hospital_id, branch_id", "appointments, encounters, lab_orders, billing_invoices, pharmacy_stock", "Only allowed branch users can access."],
    ["Patient-level table", "tenant_id, hospital_id, patient_id", "patient_identifiers, patient_addresses, patient_contacts", "Scoped to hospital patient."],
    ["Encounter-level table", "tenant_id, hospital_id, branch_id, patient_id, encounter_id", "vitals, diagnoses, prescriptions, lab_orders", "Must match the encounter branch and patient."],
]


def mermaid_blocks() -> dict[str, str]:
    return {
        "hierarchy": """
graph TD
A[Tenant / Hospital Group] --> B[Hospital]
B --> C[Branch / Facility / Location]
C --> D[Department / Healthcare Service]
C --> E[Users / Doctors / Staff]
B --> F[Patient]
F --> G[Encounter / Visit]
G --> H[Clinical Records]
G --> I[Lab Orders & Results]
G --> J[Radiology Orders & Reports]
G --> K[Prescription]
K --> L[Pharmacy Dispensing]
G --> M[Billing & Payment]
G --> N[Documents]
H --> O[FHIR Mapping Layer]
I --> O
J --> O
K --> O
M --> O
N --> O
""".strip(),
        "encounter": """
graph TD
A[Patient] --> B[Encounter]
C[Branch / Location] --> B
D[Doctor / Practitioner] --> E[Encounter Participants]
B --> E
B --> F[Vitals / Observations]
B --> G[Diagnoses / Conditions]
B --> H[Clinical Notes]
B --> I[Procedures]
B --> J[Care Plans]
B --> K[Documents]
B --> L[Lab Orders]
B --> M[Prescription]
B --> N[Billing Invoice]
""".strip(),
        "lab": """
graph LR
A[Encounter] --> B[Lab Order]
B --> C[Lab Order Items]
C --> D[Test Master]
D --> E[Test Parameters]
C --> F[Sample / Specimen]
C --> G[Result]
G --> H[Result Parameters / Observations]
G --> I[Diagnostic Report]
""".strip(),
        "reporting": """
graph LR
A[Transactional Tables] --> B[Aggregation Job]
B --> C[Summary Tables]
C --> D[Dashboard]
C --> E[Reports]
C --> F[AI Command Center]
""".strip(),
        "erd": """
erDiagram
    tenants ||--o{ hospitals : owns
    hospitals ||--o{ branches : has
    hospitals ||--o{ users : employs
    hospitals ||--o{ patients : registers
    branches ||--o{ appointments : schedules
    branches ||--o{ encounters : hosts
    patients ||--o{ appointments : books
    patients ||--o{ encounters : visits
    appointments ||--o| encounters : creates
    encounters ||--o{ vitals : records
    encounters ||--o{ diagnoses : records
    encounters ||--o{ prescriptions : orders
    prescriptions ||--o{ prescription_items : contains
    encounters ||--o{ lab_orders : requests
    lab_orders ||--o{ lab_order_items : contains
    lab_order_items ||--o{ lab_results : produces
    encounters ||--o{ billing_invoices : bills
    billing_invoices ||--o{ billing_invoice_items : contains
    billing_invoices ||--o{ payment_allocations : receives
    payments ||--o{ payment_allocations : allocates
    medicine_master ||--o{ medicine_batches : batches
    medicine_batches ||--o{ pharmacy_sale_items : sold_as
    pharmacy_sales ||--o{ pharmacy_sale_items : contains
    patients ||--o{ documents : owns
    encounters ||--o{ documents : produces
    users ||--o{ audit_logs : performs
    fhir_resource_mapping }o--|| patients : maps
    hl7_message_log ||--o{ lab_orders : may_create
""".strip(),
    }


def build_doc(diagrams: dict[str, Path]) -> Document:
    doc = setup_doc()

    # Cover Page
    cover = doc.add_table(rows=1, cols=1)
    cell = cover.cell(0, 0)
    shade(cell, "14325C")
    border(cell, "14325C")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Plasmit Hospital Management System")
    r.font.name = "Aptos Display"
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("High-Level Database Contract, Table Relationship Blueprint & FHIR Mapping")
    r2.font.name = "Aptos Display"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(*rgb("D8E8FF"))
    doc.add_paragraph()
    picture(doc, diagrams["hierarchy"], "Database contract hierarchy for tenant, hospital, branch, patient and encounter.")
    table(
        doc,
        ["Item", "Details"],
        [
            ["Project", "Plasmit Global Hospital Management System"],
            ["Document", "High-Level Database Contract & FHIR Mapping Blueprint"],
            ["Version", "1.0"],
            ["Prepared For", "Engineering, Architecture, QA, Integration and Product Teams"],
            ["Prepared By", "Plasmit Engineering Team"],
            ["Technology", "Spring Boot microservices, Java 17, MySQL, JDBC/NamedParameterJdbcTemplate, Flyway"],
        ],
    )
    doc.add_page_break()

    heading(doc, "Table of Contents", 1)
    bullets(doc, [
        "1. Executive Summary",
        "2. Core Database Contract Principles",
        "3. Recommended Multi-Tenant Hierarchy",
        "4. Core Foundation Tables and Relationships",
        "5. Patient-Centric Tables and Relationships",
        "6. Appointment Tables and Relationships",
        "7. Encounter-Centric Clinical Model",
        "8. Prescription and Medication Tables",
        "9. Lab / Diagnostics Tables and Relationships",
        "10. Radiology Tables and Relationships",
        "11. Pharmacy and Inventory Tables",
        "12. Billing, Payment and Insurance Tables",
        "13. IPD, Ward, Room and Bed Tables",
        "14. Surgery / OT Tables",
        "15. Documents and Reports",
        "16. Audit and Compliance Tables",
        "17. FHIR Integration Tables",
        "18. HL7 Integration Tables",
        "19. Reporting and Denormalized Tables",
        "20. Master Table Contract",
        "21. Complete Table Relationship Matrix",
        "22. Module Ownership Matrix",
        "23. Foreign Key Contract Rules",
        "24. Tenant Isolation and Query Contract",
        "25. Indexing, Migration and API Impact Contract",
        "26. Mermaid ERD Visual Contract",
        "27. Developer Agreement Checklist",
        "28. Final Database Contract Summary",
    ])
    doc.add_page_break()

    heading(doc, "1. Executive Summary", 1)
    para(doc, "This document is the formal high-level database contract for the Plasmit Global Hospital Management System. It defines the agreed tables, table relationships, ownership boundaries, tenant isolation rules, FHIR/HL7 mapping responsibilities, migration rules and API query rules that every developer must follow.")
    callout(doc, "Contract purpose", [
        "Create a fixed agreement before parallel backend development starts.",
        "Prevent duplicate tables, inconsistent naming and broken relationships.",
        "Protect tenant, hospital and branch data isolation.",
        "Ensure patient, encounter, lab, pharmacy, billing and document records are FHIR-ready.",
        "Ensure HL7 inbound/outbound messages are logged and traceable.",
        "Make Flyway migrations and review workflows mandatory.",
    ])

    heading(doc, "2. Core Database Contract Principles", 1)
    bullets(doc, [
        "Every major table must include tenant_id and hospital_id.",
        "Every branch-level operational table must include branch_id.",
        "Patient should be hospital-level by default.",
        "Encounter should be branch-level.",
        "Billing, lab, pharmacy, appointment, bed, ward and stock records should be branch-level.",
        "No table should be created without ownership.",
        "No duplicate table should be created for the same business entity.",
        "No backend query should fetch data without tenant/hospital/branch filters.",
        "All schema changes must use Flyway migration.",
        "Old merged migration files must never be edited.",
        "Every table must have audit columns.",
        "Every critical action must create audit log.",
        "FHIR mapping must be documented for clinical, patient, billing, lab, pharmacy and document modules.",
        "HL7 inbound/outbound messages must be logged.",
        "Reporting tables must not replace transactional normalized tables.",
    ])

    heading(doc, "3. Recommended Multi-Tenant Hierarchy", 1)
    picture(doc, diagrams["hierarchy"], "Tenant to FHIR mapping hierarchy.")
    code(doc, mermaid_blocks()["hierarchy"])

    heading(doc, "4. Core Foundation Tables and Relationships", 1)
    picture(doc, diagrams["foundation"], "Foundation relationship contract.")
    table(doc, ["Parent Table", "Child Table", "Relationship", "Join Key", "Business Meaning", "FHIR Mapping"], FOUNDATION_RELATIONSHIPS)

    heading(doc, "5. Patient-Centric Tables and Relationships", 1)
    para(doc, "Patient is hospital-level by default. A patient can visit multiple branches under the same hospital. Branch-specific activity begins from appointment and encounter.")
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["patients", "hospitals 1:N patients", "Patient"],
        ["patient_identifiers", "patients 1:N identifiers", "Patient.identifier"],
        ["patient_addresses", "patients 1:N addresses", "Patient.address"],
        ["patient_contacts", "patients 1:N contacts", "Patient.contact / RelatedPerson"],
        ["patient_allergies", "patients 1:N allergies", "AllergyIntolerance"],
        ["patient_insurance", "patients 1:N insurance policies", "Coverage"],
        ["patient_documents", "patients 1:N documents", "DocumentReference"],
        ["patient_consent_records", "patients 1:N consent records", "Consent"],
        ["patient_record_access_logs", "patients 1:N access logs", "AuditEvent"],
    ])

    heading(doc, "6. Appointment Tables and Relationships", 1)
    table(doc, ["Table", "Primary Relationship", "FHIR Mapping"], [
        ["appointment_slots", "branches/doctors 1:N slots", "Schedule / Slot if required"],
        ["appointments", "patients, branches and doctors 1:N appointments", "Appointment"],
        ["appointment_status_history", "appointments 1:N status history", "AuditEvent / Provenance"],
        ["appointment_reminders", "appointments 1:N reminders", "Communication"],
        ["appointment_cancellations", "appointments 0/1 cancellation record", "Appointment.status / AuditEvent"],
    ])

    heading(doc, "7. Encounter-Centric Clinical Model", 1)
    picture(doc, diagrams["encounter"], "Encounter-centric relationship model.")
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["encounters", "patients 1:N encounters; branches 1:N encounters; appointment 0/1:1 encounter", "Encounter"],
        ["encounter_participants", "encounters 1:N participants; users/practitioners N:1", "Encounter.participant / PractitionerRole"],
        ["encounter_status_history", "encounters 1:N status events", "Provenance / AuditEvent"],
        ["vitals", "encounters 1:N vitals", "Observation"],
        ["diagnoses", "encounters 1:N diagnoses", "Condition"],
        ["clinical_notes", "encounters 1:N clinical notes", "DocumentReference / Composition"],
        ["clinical_documents", "encounters 1:N clinical documents", "DocumentReference"],
        ["procedures", "encounters 1:N procedures", "Procedure"],
        ["care_plans", "encounters 1:N care plans", "CarePlan"],
    ])
    code(doc, mermaid_blocks()["encounter"])

    heading(doc, "8. Prescription and Medication Tables", 1)
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["prescriptions", "encounters 1:N prescriptions", "MedicationRequest"],
        ["prescription_items", "prescriptions 1:N items; medicine_master N:1", "MedicationRequest.dosageInstruction"],
        ["medicine_master", "medicine catalog used by prescription and pharmacy", "Medication"],
        ["medication_instructions", "prescription_items 1:N reusable instructions", "MedicationRequest.dosageInstruction"],
        ["medication_administration_records", "inpatient medication administration events", "MedicationAdministration"],
        ["pharmacy_dispense", "prescriptions 1:N dispense events", "MedicationDispense"],
        ["pharmacy_dispense_items", "dispense event 1:N batch/medicine items", "MedicationDispense"],
    ])

    heading(doc, "9. Lab / Diagnostics Tables and Relationships", 1)
    picture(doc, diagrams["lab"], "Lab order to report relationship contract.")
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["lab_test_master", "Master catalog for lab tests", "ObservationDefinition / Code mapping"],
        ["lab_test_parameters", "lab_test_master 1:N parameters", "Observation.component"],
        ["lab_orders", "encounters 1:N lab orders", "ServiceRequest"],
        ["lab_order_items", "lab_orders 1:N items; lab_test_master N:1", "ServiceRequest"],
        ["lab_samples", "lab_order_items 1:N samples", "Specimen"],
        ["lab_results", "lab_order_items 1:N results", "Observation"],
        ["lab_result_parameters", "lab_results 1:N result values", "Observation.component"],
        ["diagnostic_reports", "lab_results may generate reports", "DiagnosticReport"],
        ["lab_result_review_history", "lab_results 1:N reviews", "Provenance / AuditEvent"],
    ])
    code(doc, mermaid_blocks()["lab"])

    heading(doc, "10. Radiology Tables and Relationships", 1)
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["radiology_test_master", "Radiology/imaging test catalog", "ServiceRequest code"],
        ["radiology_orders", "encounters 1:N radiology orders", "ServiceRequest"],
        ["radiology_order_items", "radiology_orders 1:N order items", "ServiceRequest"],
        ["radiology_reports", "radiology_order_items generate reports", "DiagnosticReport"],
        ["imaging_studies", "radiology_reports may connect to imaging studies", "ImagingStudy"],
        ["radiology_report_review_history", "radiology_reports 1:N reviews", "Provenance / AuditEvent"],
    ])

    heading(doc, "11. Pharmacy and Inventory Tables", 1)
    picture(doc, diagrams["pharmacy"], "Pharmacy, batch, stock and dispensing relationship contract.")
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["medicine_master", "Master medicine catalog", "Medication"],
        ["medicine_categories", "medicine grouping", "Internal"],
        ["medicine_batches", "medicine_master 1:N batches", "Medication.batch"],
        ["pharmacy_stock", "branches 1:N stock; batches 1:N stock rows", "Internal"],
        ["stock_movements", "tracks every stock change", "Internal"],
        ["pharmacy_sales", "prescription may connect to pharmacy sale", "MedicationDispense"],
        ["pharmacy_sale_items", "pharmacy_sales 1:N sale items; batches N:1", "MedicationDispense"],
        ["pharmacy_returns", "sales 1:N return events", "Internal"],
        ["purchase_orders / goods_receipts", "procurement creates stock source", "Internal"],
    ])

    heading(doc, "12. Billing, Payment and Insurance Tables", 1)
    picture(doc, diagrams["billing"], "Billing, payment and insurance relationship contract.")
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["billing_invoices", "patients and encounters 1:N invoices", "Account"],
        ["billing_invoice_items", "billing_invoices 1:N charge items", "ChargeItem"],
        ["payments", "payment events", "PaymentReconciliation if required"],
        ["payment_allocations", "billing_invoices and payments N:M link", "ExplanationOfBenefit if required"],
        ["refunds", "billing_invoices 1:N refunds", "Internal / PaymentReconciliation"],
        ["insurance_policies", "patients 1:N policies", "Coverage"],
        ["insurance_claims", "billing_invoices may connect to claims", "Claim"],
        ["insurance_claim_items", "claims 1:N claim items", "Claim.item"],
        ["claim_documents", "claims 1:N documents", "DocumentReference"],
        ["claim_status_history", "claims 1:N status events", "Provenance / ExplanationOfBenefit"],
    ])

    heading(doc, "13. IPD, Ward, Room and Bed Tables", 1)
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["admissions", "encounters may create admissions", "Encounter"],
        ["wards", "branches 1:N wards", "Location"],
        ["rooms", "wards 1:N rooms", "Location"],
        ["beds", "rooms 1:N beds", "Location"],
        ["bed_allocations", "admissions 1:N bed allocations", "Encounter.location"],
        ["nursing_notes", "admissions 1:N nursing notes", "DocumentReference"],
        ["nursing_tasks", "admissions 1:N nursing tasks", "Task"],
        ["discharge_plans", "admissions 1:N plans", "CarePlan"],
        ["discharge_summaries", "admissions 1:N summaries", "Composition + DocumentReference"],
    ])

    heading(doc, "14. Surgery / OT Tables", 1)
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["operation_theatres", "branches 1:N operation theatres", "Location"],
        ["surgery_cases", "encounters 1:N surgery cases", "Procedure"],
        ["surgery_team_members", "surgery_cases 1:N team members", "PractitionerRole"],
        ["surgery_checklists", "surgery_cases 1:N checklist records", "QuestionnaireResponse if required"],
        ["anesthesia_records", "surgery_cases 1:N anesthesia records", "Observation"],
        ["surgery_notes", "surgery_cases 1:N notes", "DocumentReference"],
        ["post_operation_notes", "surgery_cases 1:N post-op notes", "DocumentReference"],
    ])

    heading(doc, "15. Documents and Reports", 1)
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["documents", "patients/encounters 1:N documents", "DocumentReference"],
        ["document_versions", "documents 1:N versions", "DocumentReference.version"],
        ["document_access_logs", "documents 1:N access logs", "AuditEvent"],
        ["document_signatures", "documents 1:N signatures", "Provenance / DocumentReference"],
    ])

    heading(doc, "16. Audit and Compliance Tables", 1)
    table(doc, ["Table", "Relationship Contract", "FHIR Mapping"], [
        ["audit_logs", "every critical module writes audit logs", "AuditEvent"],
        ["user_login_history", "users 1:N login records", "AuditEvent"],
        ["patient_record_access_logs", "patients 1:N access logs", "AuditEvent"],
        ["consent_records", "patients 1:N consent records", "Consent"],
        ["break_glass_access_logs", "emergency access log", "AuditEvent"],
        ["data_export_logs", "patient data export trace", "AuditEvent / Provenance"],
    ])

    heading(doc, "17. FHIR Integration Tables", 1)
    picture(doc, diagrams["integration"], "FHIR and HL7 integration readiness contract.")
    table(doc, ["FHIR Table", "Purpose", "Key Columns"], [
        ["fhir_resource_mapping", "Connect internal table records to FHIR resource IDs.", "id, tenant_id, hospital_id, branch_id, internal_table_name, internal_record_id, fhir_resource_type, fhir_resource_id, fhir_version, profile_url, sync_status"],
        ["fhir_resource_store", "Store generated or received FHIR JSON when required.", "id, mapping_id, resource_type, resource_id, resource_json, version, status"],
        ["fhir_api_audit_log", "Track every FHIR API call.", "id, tenant_id, hospital_id, request_method, resource_type, status_code, performed_by"],
        ["fhir_sync_status", "Track sync with external systems.", "id, mapping_id, target_system, sync_status, last_synced_at"],
        ["fhir_profile_registry", "Register supported FHIR profiles.", "id, profile_url, resource_type, version, status"],
        ["fhir_code_system_mapping", "Map internal codes to SNOMED, LOINC, ICD, RxNorm.", "id, internal_code, standard_system, standard_code"],
        ["fhir_value_set_mapping", "Manage allowed FHIR values.", "id, value_set_url, internal_value, fhir_value"],
    ])
    table(doc, ["HMS Concept", "Internal Tables", "FHIR Resource", "Integration Purpose", "Owner"], FHIR_MAPPING)

    heading(doc, "18. HL7 Integration Tables", 1)
    table(doc, ["HL7 Table", "Purpose"], [
        ["hl7_message_log", "Store inbound/outbound raw HL7 message, direction, type, status and processing trace."],
        ["hl7_message_error_log", "Store parse, validation and mapping errors."],
        ["hl7_external_identifier_mapping", "Map external system identifiers to internal records."],
        ["hl7_patient_mapping", "External patient ID to internal patient ID."],
        ["hl7_order_mapping", "External order ID to internal lab/radiology order."],
        ["hl7_result_mapping", "External result ID to internal observation/result."],
    ])
    table(doc, ["HL7 Message Type", "Purpose"], [
        ["ADT", "Admission, discharge and transfer messages."],
        ["ORM", "Order messages, commonly lab/radiology orders."],
        ["ORU", "Observation/result messages."],
        ["SIU", "Scheduling messages."],
        ["DFT", "Financial transaction messages."],
        ["MDM", "Medical document management messages."],
    ])

    heading(doc, "19. Reporting and Denormalized Tables", 1)
    picture(doc, diagrams["reporting"], "Transactional to denormalized reporting contract.")
    table(doc, ["Summary Table", "Source Tables", "Purpose"], [
        ["daily_branch_revenue_summary", "billing_invoices, payments", "Branch revenue dashboard."],
        ["daily_patient_visit_summary", "appointments, encounters", "Visit volume dashboard."],
        ["doctor_performance_summary", "encounters, orders, billing", "Doctor productivity analytics."],
        ["department_collection_summary", "billing_invoice_items, departments", "Department revenue tracking."],
        ["lab_test_volume_summary", "lab_orders, lab_results", "Lab workload analytics."],
        ["pharmacy_stock_snapshot", "pharmacy_stock, stock_movements", "Stock and expiry dashboard."],
        ["bed_occupancy_summary", "admissions, beds, bed_allocations", "IPD occupancy dashboard."],
        ["patient_journey_summary", "encounter, lab, pharmacy, billing, discharge", "AI command center and operational tracking."],
    ])
    code(doc, mermaid_blocks()["reporting"])

    heading(doc, "20. Master Table Contract", 1)
    table(doc, ["Master Type", "Examples", "Scope Rule"], [
        ["Global standard masters", "gender_master, blood_group_master, country_master, state_master, city_master", "No hospital_id unless customization is required."],
        ["Workflow masters", "encounter_type_master, appointment_status_master, invoice_status_master, payment_mode_master", "Global or tenant-level depending on product design."],
        ["FHIR/HL7 masters", "fhir_resource_type_master, hl7_message_type_master", "Global standard masters."],
        ["Clinical code masters", "diagnosis_code_master, loinc_code_master, snomed_code_master, icd10_code_master, rxnorm_code_master", "Global standard, with hospital extension table if customized."],
        ["Hospital-specific masters", "service_catalog, charge_master, medicine_master", "Must include tenant_id and hospital_id."],
        ["Branch-specific masters", "counters, rooms, beds, pharmacy_stock", "Must include tenant_id, hospital_id and branch_id."],
    ])

    heading(doc, "21. Complete Table Relationship Matrix", 1)
    table(doc, ["Module", "Table Name", "Primary Key", "Foreign Keys", "Connected Parent Table", "Connected Child Table", "Relationship Type", "Tenant Scope", "Branch Scope", "FHIR Resource", "Owner Team"], RELATIONSHIP_MATRIX)

    heading(doc, "22. Module Ownership Matrix", 1)
    table(doc, ["Module", "Owner Developer/Team", "Tables Owned", "Shared Tables Used", "FHIR Resource Owner", "Review Required From"], OWNERSHIP)

    heading(doc, "23. Foreign Key Contract Rules", 1)
    bullets(doc, [
        "Use foreign keys for core transactional integrity where safe.",
        "Avoid unnecessary cross-microservice hard coupling if services are fully independent.",
        "Use indexed foreign key columns.",
        "Do not cascade delete patient, encounter, billing or clinical records.",
        "Use soft delete with is_deleted.",
        "For audit and logs, avoid cascade delete.",
        "Use status columns for lifecycle management.",
        "patients.id connects to encounters.patient_id.",
        "encounters.id connects to lab_orders.encounter_id, prescriptions.encounter_id and billing_invoices.encounter_id.",
        "branches.id connects to appointments.branch_id, encounters.branch_id and pharmacy_stock.branch_id.",
        "users.id connects to created_by, updated_by, doctor_id and performed_by.",
        "billing_invoices.id connects to payment_allocations.invoice_id.",
        "lab_order_items.id connects to lab_results.lab_order_item_id.",
    ])

    heading(doc, "24. Tenant Isolation and Query Contract", 1)
    table(doc, ["Table Type", "Required Columns", "Example Tables", "Access Rule"], TENANT_SCOPE)
    heading(doc, "Standard Query Contract", 2)
    code(doc, """
SELECT *
FROM patients
WHERE tenant_id = :tenantId
  AND hospital_id = :hospitalId
  AND is_deleted = 0;

SELECT *
FROM appointments
WHERE tenant_id = :tenantId
  AND hospital_id = :hospitalId
  AND branch_id IN (:allowedBranchIds)
  AND is_deleted = 0;

SELECT *
FROM vitals
WHERE tenant_id = :tenantId
  AND hospital_id = :hospitalId
  AND branch_id = :branchId
  AND patient_id = :patientId
  AND encounter_id = :encounterId
  AND is_deleted = 0;
""")

    heading(doc, "25. Indexing, Migration and API Impact Contract", 1)
    table(doc, ["Contract Area", "Rule"], [
        ["Audit columns", "Every major table must include id, tenant_id, hospital_id, branch_id where applicable, created_by, updated_by, created_at, updated_at, is_deleted and status where applicable."],
        ["Money columns", "Store amount in paise/cents as BIGINT. Do not use FLOAT/DOUBLE for money. Use DECIMAL only where legally required."],
        ["Tenant index", "Every tenant scoped table must have composite tenant/hospital index."],
        ["Branch/date/status index", "Every branch transaction table must have branch/date/status index."],
        ["Patient index", "Every patient activity table must have patient_id index."],
        ["Encounter index", "Every encounter-level table must have encounter_id index."],
        ["Integration index", "Every external integration table must have external identifier index."],
        ["Migration", "Flyway must be used. Use timestamp-based migration names. Never edit already merged migration."],
        ["API impact", "Every API must mention table used, relationship dependency, tenant scope, branch scope, FHIR impact and audit impact."],
    ])
    code(doc, """
CREATE INDEX idx_patients_tenant_hospital_mrn
ON patients (tenant_id, hospital_id, mrn);

CREATE INDEX idx_encounters_branch_patient
ON encounters (tenant_id, hospital_id, branch_id, patient_id);

CREATE INDEX idx_lab_orders_branch_status
ON lab_orders (tenant_id, hospital_id, branch_id, status, ordered_at);

V202606031000__create_patient_tables.sql
V202606031030__create_encounter_tables.sql
V202606031100__create_lab_tables.sql
""")
    table(doc, ["API Example", "Contract"], [
        ["Create Lab Order API - Tables", "lab_orders, lab_order_items, audit_logs"],
        ["Create Lab Order API - FHIR Impact", "Creates ServiceRequest mapping in fhir_resource_mapping."],
        ["Create Lab Order API - Tenant Scope", "tenant_id, hospital_id and branch_id must come from token/TenantContext."],
        ["Create Lab Order API - Audit", "Must write LAB_ORDER_CREATED audit event."],
    ])

    heading(doc, "26. Mermaid ERD Visual Contract", 1)
    code(doc, mermaid_blocks()["erd"])

    heading(doc, "27. Developer Agreement Checklist", 1)
    bullets(doc, [
        "Is this table already existing?",
        "Which module owns it?",
        "Is it global, tenant, hospital, branch, patient or encounter level?",
        "Does it need tenant_id, hospital_id or branch_id?",
        "Does it need patient_id?",
        "Does it need encounter_id?",
        "Does it map to FHIR?",
        "Does it need HL7 integration?",
        "Does it need audit log?",
        "Does it need reporting summary?",
        "Which other table connects with it?",
        "Which API will use it?",
        "Which index is required?",
        "Has Flyway migration been created?",
        "Has database architect or module owner reviewed the table?",
    ])

    heading(doc, "28. Final Database Contract Summary", 1)
    callout(doc, "Final contract recommendation", [
        "Foundation tables are created first.",
        "Patient is the central hospital-level entity.",
        "Encounter is the central branch-level clinical entity.",
        "Lab, pharmacy, radiology, prescription and billing connect mostly through encounter.",
        "FHIR layer maps internal records to external healthcare resources.",
        "HL7 layer logs and maps inbound/outbound legacy messages.",
        "Reporting tables are generated from normalized source tables.",
        "Tenant, hospital and branch isolation is mandatory.",
        "No developer can create or modify shared tables without architect review.",
        "This document is the baseline contract for database development.",
    ], bg="E8F7F3", title_color="0F766E")

    return doc


def md_table(headers: list[str], rows: list[list[str]]) -> list[str]:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        safe = [cell.replace("\n", "<br>").replace("|", "/") for cell in row]
        lines.append("| " + " | ".join(safe) + " |")
    return lines


def build_markdown(diagrams: dict[str, Path]) -> str:
    blocks = mermaid_blocks()
    lines = [
        "# Plasmit Hospital Management System - High-Level Database Contract, Table Relationship Blueprint & FHIR Mapping",
        "",
        "## Executive Summary",
        "",
        "This document is the database contract between database architect, backend developers, frontend developers, FHIR/HL7 integration developers, QA and project leadership. It defines what tables will exist, how they connect, which module owns them, and what rules every API and migration must follow.",
        "",
        "## Contract Principles",
        "",
        "- Every major table must include `tenant_id` and `hospital_id`.",
        "- Every branch-level operational table must include `branch_id`.",
        "- Patient is hospital-level by default.",
        "- Encounter is branch-level and becomes the clinical anchor.",
        "- No table can be created without owner and review.",
        "- No backend query can fetch data without tenant/hospital/branch filters.",
        "- Flyway migrations are mandatory.",
        "- FHIR mapping must be documented for patient, clinical, lab, pharmacy, billing and documents.",
        "- HL7 inbound/outbound messages must be logged.",
        "- Reporting tables do not replace normalized transaction tables.",
        "",
        "## Diagram Assets",
        "",
    ]
    for key, path in diagrams.items():
        lines.append(f"- {key}: `{path.relative_to(ROOT).as_posix()}`")
    lines.extend([
        "",
        "## Multi-Tenant Hierarchy",
        "",
        "```mermaid",
        blocks["hierarchy"],
        "```",
        "",
        "## Foundation Relationships",
        "",
    ])
    lines.extend(md_table(["Parent Table", "Child Table", "Relationship", "Join Key", "Business Meaning", "FHIR Mapping"], FOUNDATION_RELATIONSHIPS))
    lines.extend(["", "## Module-Wise Table List", ""])
    lines.extend(md_table(["Module", "Tables", "FHIR Mapping"], MODULE_TABLES))
    lines.extend(["", "## Complete Table Relationship Matrix", ""])
    lines.extend(md_table(["Module", "Table Name", "Primary Key", "Foreign Keys", "Connected Parent Table", "Connected Child Table", "Relationship Type", "Tenant Scope", "Branch Scope", "FHIR Resource", "Owner Team"], RELATIONSHIP_MATRIX))
    lines.extend(["", "## FHIR Mapping Matrix", ""])
    lines.extend(md_table(["HMS Concept", "Internal Tables", "FHIR Resource", "Integration Purpose", "Owner"], FHIR_MAPPING))
    lines.extend(["", "## Module Ownership Matrix", ""])
    lines.extend(md_table(["Module", "Owner Developer/Team", "Tables Owned", "Shared Tables Used", "FHIR Resource Owner", "Review Required From"], OWNERSHIP))
    lines.extend([
        "",
        "## Tenant Isolation Contract",
        "",
    ])
    lines.extend(md_table(["Table Type", "Required Columns", "Example Tables", "Access Rule"], TENANT_SCOPE))
    lines.extend([
        "",
        "## Standard Query Contract",
        "",
        "```sql",
        "SELECT *",
        "FROM patients",
        "WHERE tenant_id = :tenantId",
        "  AND hospital_id = :hospitalId",
        "  AND is_deleted = 0;",
        "",
        "SELECT *",
        "FROM appointments",
        "WHERE tenant_id = :tenantId",
        "  AND hospital_id = :hospitalId",
        "  AND branch_id IN (:allowedBranchIds)",
        "  AND is_deleted = 0;",
        "```",
        "",
        "## Encounter Mermaid Diagram",
        "",
        "```mermaid",
        blocks["encounter"],
        "```",
        "",
        "## Lab Mermaid Diagram",
        "",
        "```mermaid",
        blocks["lab"],
        "```",
        "",
        "## Reporting Mermaid Diagram",
        "",
        "```mermaid",
        blocks["reporting"],
        "```",
        "",
        "## High-Level ERD",
        "",
        "```mermaid",
        blocks["erd"],
        "```",
        "",
        "## Developer Agreement Checklist",
        "",
        "- Is this table already existing?",
        "- Which module owns it?",
        "- Is it global, tenant, hospital, branch, patient or encounter level?",
        "- Does it need tenant_id/hospital_id/branch_id?",
        "- Does it need patient_id or encounter_id?",
        "- Does it map to FHIR?",
        "- Does it need HL7 integration?",
        "- Does it need audit log or reporting summary?",
        "- Which other table connects with it?",
        "- Which API will use it?",
        "- Which index is required?",
        "",
        "## Final Recommendation",
        "",
        "Foundation tables are created first. Patient is the central hospital-level entity. Encounter is the central branch-level clinical entity. Lab, pharmacy, radiology, prescription and billing connect mostly through encounter. FHIR and HL7 are integration layers. Reporting tables are generated from normalized source tables. Tenant/hospital/branch isolation is mandatory.",
    ])
    return "\n".join(lines)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    diagrams = create_diagrams()
    doc = build_doc(diagrams)
    doc.save(DOCX)
    MD.write_text(build_markdown(diagrams), encoding="utf-8")
    print(f"created: {DOCX}")
    print(f"created: {MD}")
    print(f"diagrams: {len(diagrams)}")


if __name__ == "__main__":
    main()
