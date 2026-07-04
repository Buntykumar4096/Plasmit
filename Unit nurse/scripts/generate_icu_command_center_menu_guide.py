from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Plasmit-HMS-ICU-Command-Center-Menu-and-Screen-Guide.docx"

BLUE = "078BCB"
NAVY = "17365D"
LIGHT_BLUE = "EAF5FB"
LIGHT_GREY = "F4F7FA"
WHITE = "FFFFFF"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(75, 85, 99)


MENUS = [
    (
        "1. Command",
        "Shows the current ICU position and the work that needs immediate attention.",
        [
            ("Command Center", "Shows the current summary and priority of all ICU patients.", "One ICU patient.", "Patient, Diagnosis, Risk, Ventilation, Intake / Output, Medication, Events"),
            ("Executive Dashboard", "Compares occupancy, workload and readiness across ICU units.", "One ICU unit.", "ICU Unit, Occupancy, Critical, Ventilator, Alerts, Medication, Documentation, Device, Transfer, Owner, Action"),
            ("Notifications & Tasks", "Lists clinical alerts and pending tasks in priority order.", "One notification or task.", "Patient / Notification, Source, Priority, SLA, Owner, Status, Action"),
        ],
    ),
    (
        "2. Patients",
        "Supports patient search, bed review, admission and discharge workflows.",
        [
            ("Patient Search", "Finds a patient by name, ID, bed or ICU unit.", "One patient.", "Patient Identity, Bed / Unit, Doctor, Risk, Ventilator, Alerts, Status, Action"),
            ("Smart Bed View", "Shows the patient, devices and care status for each ICU bed.", "One ICU bed.", "Bed Context, Patient, Ventilation, Device Work, Care Team, Flow, Action"),
            ("Admissions", "Manages ICU admission requests, bed selection and team readiness.", "One admission request.", "Patient, ICU Need, Unit, Bed, Readiness, Doctor, Nurse, Status, Action"),
            ("Discharges", "Tracks discharge or transfer readiness and bed release.", "One discharge or transfer case.", "Patient, ICU Context, Destination, Readiness, Blockers, Clearance, Handover, Transport, Bed Release, Action"),
        ],
    ),
    (
        "3. Critical Care",
        "Manages ICU operations, devices, alerts, rounds and escalations.",
        [
            ("ICU Operations", "Shows bed readiness, staffing and operational blockers.", "One bed or unit item.", "Bed / Patient, Status, Staffing, Devices, Blocker, SLA, Action"),
            ("Device Monitoring", "Monitors the clinical devices connected to each bed.", "One bed device set.", "Bed / Patient, Monitor, Ventilator, Pump, Gateway, Signal, Last Data, Issue, Assigned To, Action"),
            ("Clinical Alerts", "Tracks patient alerts by severity, source and response time.", "One clinical alert.", "Patient, Trigger, Classification, Severity, Source, Assigned To, SLA, Route / Action"),
            ("ICU Rounds", "Shows patient priority and pending work before doctor rounds.", "One ICU patient.", "Patient / Bed, Risk, Latest Vitals, Ventilation, Alerts, Due Medicines, Pending Labs, Tasks, Round Entry"),
            ("Escalation Center", "Manages urgent cases, assigned staff and response time.", "One patient escalation.", "Patient, Priority, Trigger, Source, Assigned To, SLA, Status, Review"),
        ],
    ),
    (
        "4. Clinical Workspace",
        "Contains the selected patient's summary, notes, orders and family communication.",
        [
            ("Patient Overview", "Shows the selected patient's current clinical summary.", "One selected patient workspace.", "Patient & Team, Latest Observation, Medication / Infusions, Investigations, Active Issues, Timeline"),
            ("Progress Notes", "Records and reviews notes entered by the clinical team.", "One clinical note.", "Date / Time, Author, Role / Type, Note Summary, Status, Action"),
            ("Orders & Care Plans", "Keeps doctor orders and nursing care plans in one workflow.", "One order or care task.", "Patient Context, Order / Care Plan, Priority, Assigned To, Due, Status, Action"),
            ("Family Communication", "Records family updates, consent discussions and follow-ups.", "One family communication entry.", "Patient, Bed, Communication Type, Attendees, Summary, Questions, Consent, Follow-up, Recorded By, Priority, Action"),
        ],
    ),
    (
        "5. Nursing",
        "Supports nursing observations, medication, handover, tasks and assessments.",
        [
            ("Nursing Station", "Shows assigned patients and current nursing workload.", "One assigned patient.", "Patient, Coverage, Workload, Medication, Observations, Tasks, Handover, Action"),
            ("Nurse Entry", "Records patient vitals and bedside observations.", "One saved observation entry.", "Patient, Date / Time, RR, SpO2, Oxygen, BP, Pulse, Temperature, GCS, Pain, Urine, Notes"),
            ("Medication Administration", "Records a scheduled dose as given, held or missed after verification.", "One scheduled medicine dose.", "Patient, Medicine, Dose, Route, Due Time, Pharmacy, Verification, Administration Status, Nurse, Action"),
            ("Patient Medication Chart", "Shows medicine history and administration status for a patient.", "One medicine order.", "Medicine, Dose, Route, Frequency, Date / Time Slots, Given / Held / Missed, Verification, Remarks"),
            ("Shift Handover", "Transfers completed and pending work between outgoing and incoming nurses.", "One nursing shift.", "Outgoing Nurse, Incoming Nurse, Reviewer, Completed Work, Pending Work, Critical Watch, Active Issues, To-do"),
            ("Tasks & Assessments", "Manages nursing tasks, assessments, risks and follow-ups.", "One task or assessment.", "Patient, Type, Priority / Risk, Source, Assigned To, Due, Status, Action"),
        ],
    ),
    (
        "6. Diagnostics",
        "Supports diagnostic orders, reports and uploaded report review.",
        [
            ("Diagnostics Hub", "Shows patient diagnostic orders and reports in one place.", "One diagnostic order or report.", "Patient, Test / Study, Category, Order / Sample Time, Status, Critical Flag, Review Status, Action"),
            ("Report Upload & Extract", "Verifies values extracted from an uploaded report.", "One extracted investigation value.", "Group, Investigation, Extracted Value, Unit, Reference Range, Flag, Confidence, Source, Review"),
        ],
    ),
    (
        "7. Tele ICU",
        "Supports remote ICU monitoring, specialist consultation and escalated cases.",
        [
            ("Remote Command Center", "Shows patient readiness and clinical status to the remote team.", "One ICU patient.", "Patient, Readiness, Vitals, Diagnostics, Ventilation, Local Team, Remote MD, SLA, Action"),
            ("Remote Consultations", "Manages specialist consultation requests and responses.", "One remote consultation.", "Patient, Consult Reason, Specialty, Readiness, Documents, Vitals, Remote MD, Status, Action"),
            ("Escalated Cases", "Tracks the current action and outcome of Tele ICU escalations.", "One escalated case.", "Patient / Case, Trigger, Severity, Source, Owner Chain, SLA, Current Action, Outcome"),
        ],
    ),
    (
        "8. Device Operations",
        "Manages ICU device mapping, connectivity and signal quality.",
        [
            ("Edge Device Management", "Manages connected devices and their current health.", "One device.", "Bed / Patient, Device ID, Type, Connectivity, Signal, Last Data, Issue, Assigned To, Action"),
            ("Device Mapping", "Maps a device to the correct bed and patient.", "One bed-patient mapping.", "Bed, Patient, Monitor, Ventilator, Pump, Gateway, Mapping Status, Action"),
            ("Connectivity Dashboard", "Monitors network and gateway connection status.", "One gateway or device connection.", "Gateway / Device, Unit, Network Status, Uptime, Last Data, Issue, Assigned To, Action"),
            ("Signal Health", "Shows signal strength, delay and missing clinical data.", "One device signal stream.", "Bed / Patient, Device, Signal Strength, Delay, Last Data, Missing Data, Issue, Action"),
        ],
    ),
    (
        "9. Clinical Intelligence",
        "Shows patient risk and early deterioration with the contributing factors.",
        [
            ("Patient Risk Center", "Shows high-risk patients and the main reasons for their risk.", "One ICU patient.", "Patient, Risk Score, Vitals, Ventilation, Infection, Medication, Device, Tasks, Action"),
            ("Early Warning Scores", "Shows the early warning score and trend calculated from observations.", "One patient score.", "Patient, Score, Contributing Vitals, Risk Band, Trend, Observation Frequency, Trigger, Action"),
        ],
    ),
    (
        "10. Analytics",
        "Shows trends in ICU performance, clinical outcomes, devices and system usage.",
        [
            ("Operational Analytics", "Shows occupancy, demand, workload and response performance.", "One ICU unit or time period.", "Unit, Occupancy, Demand, Length of Stay, Workload, Alerts, Response Time, Trend"),
            ("Clinical Analytics", "Compares clinical quality and outcome indicators.", "One clinical KPI.", "Metric, Unit, Period, Current Value, Target, Trend, Status"),
            ("Device Analytics", "Shows device use, uptime, downtime and issues.", "One device or unit KPI.", "Device / Unit, Utilization, Uptime, Issues, Downtime, Response Time, Assigned Team"),
            ("Pilot Outcome Dashboard", "Compares pilot baseline, current result and target.", "One pilot KPI.", "KPI, Baseline, Current, Target, Change, Status"),
            ("Adoption Analytics", "Shows actual system use by module and role.", "One module-role record.", "Module, Role, Usage, Key Actions, Completion, Status"),
        ],
    ),
    (
        "11. Administration",
        "Manages users, permissions, system settings and activity history.",
        [
            ("Users & Roles", "Manages user access, role and unit scope.", "One user or role.", "User, Role, Unit / Scope, Permissions, Status, Action"),
            ("Configuration", "Manages ICU rules, thresholds and system settings.", "One configuration rule.", "Category, Setting, Value, Scope, Status, Action"),
            ("Audit Logs", "Keeps a trace of activity performed in the system.", "One recorded activity.", "Date / Time, User, Role, Module, Patient / Record, Action, Old / New Value, IP / Device, Result"),
        ],
    ),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="D8E2EC", size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def style_document(document):
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    for style_name, size, color in (("Title", 24, NAVY), ("Heading 1", 16, NAVY), ("Heading 2", 12, BLUE)):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    header = section.header.paragraphs[0]
    header.text = "Plasmit Hospital HMS  |  ICU Command Center"
    header.style = styles["Normal"]
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = MUTED
    add_page_number(section.footer.paragraphs[0])


def add_menu_table(document, rows):
    table = document.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.45)
    table.columns[1].width = Inches(2.15)
    table.columns[2].width = Inches(1.35)
    table.columns[3].width = Inches(4.0)
    table.columns[4].width = Inches(1.4)
    headers = ("SUB-MENU", "USE", "EACH ROW", "MAIN COLUMNS / SECTIONS", "REVIEW NOTES")
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for index, text in enumerate(headers):
        cell = header_row.cells[index]
        cell.width = table.columns[index].width
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell, top=55, bottom=55)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.bold = True
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string(WHITE)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate((*values, "")):
            cell = row.cells[col_index]
            cell.width = table.columns[col_index].width
            set_cell_margins(cell, top=55, bottom=55)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, LIGHT_GREY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.size = Pt(7.7)
            if col_index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(NAVY)
    set_table_borders(table)
    return table


def build_document():
    document = Document()
    style_document(document)
    properties = document.core_properties
    properties.title = "ICU Command Center - Menu and Screen Guide"
    properties.subject = "Main menu, sub-menu, row and column reference"
    properties.author = "Plasmit Hospital HMS"
    properties.keywords = "ICU, menu, screen, row, column"

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    title.add_run("ICU Command Center - Menu and Screen Guide")

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(7)
    run = intro.add_run("Editable reference for menu purpose, screen rows and screen columns.")
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED

    for menu_index, (menu_name, menu_use, submenus) in enumerate(MENUS):
        heading = document.add_paragraph(style="Heading 1")
        heading.paragraph_format.space_before = Pt(5 if menu_index else 0)
        heading.paragraph_format.space_after = Pt(2)
        heading.add_run(menu_name)

        use_table = document.add_table(rows=1, cols=1)
        use_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        use_cell = use_table.cell(0, 0)
        set_cell_shading(use_cell, LIGHT_BLUE)
        set_cell_margins(use_cell, top=55, bottom=55)
        use_paragraph = use_cell.paragraphs[0]
        use_paragraph.paragraph_format.space_after = Pt(0)
        use_label = use_paragraph.add_run("Main menu use: ")
        use_label.bold = True
        use_label.font.color.rgb = RGBColor.from_string(NAVY)
        use_text = use_paragraph.add_run(menu_use)
        use_text.font.color.rgb = TEXT
        set_table_borders(use_table, color="C8DEEC")

        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)
        spacer.paragraph_format.line_spacing = 0.35
        add_menu_table(document, submenus)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    OUTPUT.chmod(0o666)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
